#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Offline regression tests for bridge.py; no model, Zotero, blog, or MCP calls."""
from __future__ import annotations

import importlib.util
import json
import os
import sys
import tempfile
import time
import unittest
from types import SimpleNamespace
from unittest.mock import patch
from pathlib import Path

HERE = Path(__file__).resolve().parent.parent
spec = importlib.util.spec_from_file_location("dsh_literature_bridge", HERE / "bridge.py")
assert spec and spec.loader
bridge = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = bridge
spec.loader.exec_module(bridge)

WORKFLOW = Path(os.environ.get(
    "LITERATURE_WORKFLOW_ROOT",
    str(Path.home() / "Downloads" / "文献积累工作流"),
)).resolve()


def sample_record() -> dict:
    story = "一位研究者想知道企业为何会更愿意投资新技术。她把政策出台前后的企业放在一起比较，" \
            "再观察受到政策影响和没有受到影响的企业有什么不同。这个比较像在两条相似的街道上，" \
            "一条街得到新路灯，另一条没有；如果亮灯后只有第一条街的店铺明显增加，就能进一步讨论政策的作用。" \
            "研究还检查了不同地区、企业规模和时期，尽量排除其他解释。故事中的比较不是凭空猜测，" \
            "而是依赖论文披露的数据和检验。结果仍要结合样本范围理解，不能把一次研究直接说成所有企业都一样。"
    story += story[:130]
    return {
        "meta": {"title": "A study of innovation", "authors_full": "A. Author", "first_author_surname": "Author", "year": 2020},
        "apa_citation": "Author, A. (2020). A study of innovation. Journal of Tests, 1(1), 1-10.",
        "authors": "A. Author", "year": "2020", "abstract": "摘要内容" * 50,
        "research_questions": ["RQ1：政策是否影响企业创新？"],
        "hypotheses": ["H1（对应RQ1）：政策提高创新投入。"],
        "core_theory": "创新理论与企业投资理论。",
        "contribution": "1. 补充现有创新研究。\n2. 提供新的识别证据。\n3. 讨论适用边界。",
        "data": "研究使用企业层面数据，并比较政策影响组与对照组。",
        "limitations": ["L1：样本范围有限。", "L2：部分机制仍待检验。", "L3：外部效度有限。"],
        "future_research": ["F1：扩展到其他地区。", "F2：补充长期数据。"],
        "fun_story": story,
        "confusions": ["C1：识别假设。", "C2：样本筛选。", "C3：标准误。"],
        "why_top": "研究问题重要，识别设计清晰，结论边界有讨论。",
        "core_variables": ["创新投入"], "economic_concepts": ["技术进步"],
        "economic_theories": ["企业投资理论"],
        "design": {
            "model_equation_latex": "Y_{it}=\\beta D_{it}+\\epsilon_{it}",
            "equation_label": "(1)", "variables": [], "fixed_effects": "【待补】",
            "cluster_level": "【待补】", "identification": {"strategy": "DID"},
            "mechanism": [], "heterogeneity": [], "robustness": [], "key_coefficients": [],
        },
        "claims": [{
            "field": "data", "value": "研究比较企业层面数据中的政策影响组与对照组。",
            "evidence": [{"field": "data", "page": 1,
                          "quote": "The study uses firm-level data and compares treated firms with a control group."}],
        }],
        "evidence": [{"field": "data", "page": 1,
                      "quote": "The study uses firm-level data and compares treated firms with a control group."}],
    }


class BridgeTests(unittest.TestCase):
    def test_stage_alias_and_allowlist(self):
        self.assertEqual(bridge._stage_list({"selectedStages": ["excel", "sheet"]}), ["excel", "sheet"])
        self.assertEqual(bridge._stage_list({"stages": ["design"], "selectedStages": ["lecture"]}), ["design"])
        with self.assertRaises(bridge.BridgeError):
            bridge._stage_list({"stages": ["delete"]})

    def test_evidence_and_numeric_result_gate(self):
        record = sample_record()
        result = bridge.validate_record(record, {
            1: "The study uses firm-level data and compares treated firms with a control group."
        }, WORKFLOW)
        self.assertTrue(result["ok"], result["errors"])

        record["design"]["key_coefficients"] = [{"term": "D", "coef": "0.123", "se": "【待核对】"}]
        result = bridge.validate_record(record, {
            1: "The study uses firm-level data and compares treated firms with a control group."
        }, WORKFLOW)
        self.assertFalse(result["ok"])
        self.assertTrue(any("数值结果" in e and "页码证据" in e for e in result["errors"]), result["errors"])
        self.assertTrue(any("待核对" in w for w in result["warnings"]))

    def test_evidence_quote_must_match_page(self):
        record = sample_record()
        result = bridge.validate_record(record, {1: "different source text"}, WORKFLOW)
        self.assertFalse(result["ok"])
        self.assertTrue(any("quote 未在第 1 页" in e for e in result["errors"]))

    def test_immutable_outputs_reuse_identical_and_reject_conflict(self):
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "out.txt"
            self.assertFalse(bridge._persist_immutable([(p, b"a")][0:1])[str(p)])
            self.assertTrue(bridge._persist_immutable([(p, b"a")])[str(p)])
            with self.assertRaises(bridge.BridgeError):
                bridge._persist_immutable([(p, b"b")])
            self.assertEqual(p.read_bytes(), b"a")

    @unittest.skipUnless(WORKFLOW.is_dir(), "audited workflow directory is unavailable")
    def test_prepare_imports_converter_without_zotero_side_effect(self):
        import pymupdf

        with tempfile.TemporaryDirectory() as td:
            td_path = Path(td)
            pdf = td_path / "sample.pdf"
            doc = pymupdf.open()
            page = doc.new_page()
            page.insert_text((72, 72), "A Clear First Page Title")
            page.insert_text((72, 100), "Published in 2020. DOI 10.1234/test.paper")
            doc.save(pdf)
            doc.close()
            job = td_path / "job"
            result = bridge._action_prepare({
                "workflowRoot": str(WORKFLOW), "pdfPath": str(pdf), "jobDir": str(job), "engine": "blocks",
            })
            self.assertEqual(result["status"], "ready")
            self.assertEqual(result["pageCount"], 1)
            self.assertEqual(result["year"], 2020)
            self.assertEqual(result["doi"], "10.1234/test.paper")
            self.assertTrue((job / "full.md").is_file())
            self.assertIn("0001:", (job / "pages" / "page-0001.txt").read_text(encoding="utf-8"))
            self.assertEqual(result["pagesPath"], str(job / "pages.json"))
            self.assertEqual(result["extractSchemaPath"], str(WORKFLOW / "templates" / "extract_schema.json"))
            pages = json.loads((job / "pages.json").read_text(encoding="utf-8"))
            self.assertEqual(pages[0]["page"], 1)
            self.assertIn("A Clear First Page Title", pages[0]["text"])
            self.assertFalse(any("zotero" in p.name.casefold() for p in job.iterdir()))

    @unittest.skipUnless(WORKFLOW.is_dir(), "audited workflow directory is unavailable")
    def test_render_only_selected_excel_and_preview_all_21_columns(self):
        with tempfile.TemporaryDirectory() as td:
            td_path = Path(td)
            extract = td_path / "extract.json"
            extract.write_text(json.dumps(sample_record(), ensure_ascii=False), encoding="utf-8")
            job = td_path / "job"
            job.mkdir()
            page_payload = [{"page": 1, "text": "The study uses firm-level data and compares treated firms with a control group."}]
            pages_path = job / "pages.json"
            pages_path.write_text(json.dumps(page_payload), encoding="utf-8")
            ctx = {"workflowRoot": str(WORKFLOW), "jobDir": str(job), "extractPath": str(extract),
                   "stages": ["excel"], "pagesPath": str(pages_path), "name": "Reader", "date": "2026.10.19"}
            result = bridge._render(ctx)
            self.assertEqual(result["status"], "rendered", result)
            self.assertEqual([a["stage"] for a in result["artifacts"]], ["excel"])
            self.assertEqual(result["artifacts"][0]["columns"], 21)
            preview = bridge._action_publish_preview({**ctx, "stages": ["sheet"]})
            self.assertTrue(preview["ok"], preview)
            self.assertEqual(preview["stage"], "sheet")
            row = Path(preview["tsvPath"]).read_text(encoding="utf-8").rstrip("\n")
            self.assertEqual(len(row.split("\t")), 21)
            registration = json.loads(Path(preview["registrationPath"]).read_text(encoding="utf-8"))
            self.assertEqual(registration["columnCount"], 21)
            self.assertEqual(len(registration["row"]), 21)
            self.assertFalse((job / "artifacts" / "PPT_Author_2020.html").exists())

    @unittest.skipUnless(WORKFLOW.is_dir(), "audited workflow directory is unavailable")
    def test_excel_uses_receipted_and_selected_future_html_paths(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            job = root / "job"
            job.mkdir()
            extract = root / "extract.json"
            extract.write_text(json.dumps(sample_record(), ensure_ascii=False), encoding="utf-8")
            pages = job / "pages.json"
            pages.write_text(json.dumps([{"page": 1, "text": "The study uses firm-level data and compares treated firms with a control group."}]), encoding="utf-8")
            base = {"workflowRoot": str(WORKFLOW), "jobDir": str(job), "extractPath": str(extract),
                    "pagesPath": str(pages), "name": "Reader", "date": "2026.10.19"}

            design = bridge._render({**base, "stages": ["design"],
                                     "pipelineStages": ["import", "prepare", "read", "design", "excel"]})
            self.assertEqual(design["status"], "rendered", design)
            excel = bridge._render({**base, "stages": ["excel"],
                                    "pipelineStages": ["import", "prepare", "read", "design", "excel"]})
            book = openpyxl.load_workbook(excel["artifacts"][0]["path"], data_only=True)
            self.assertEqual(book.active.cell(2, 16).value, design["artifacts"][0]["path"])
            self.assertIsNone(book.active.cell(2, 17).value)
            book.close()

        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            job = root / "job"
            job.mkdir()
            extract = root / "extract.json"
            extract.write_text(json.dumps(sample_record(), ensure_ascii=False), encoding="utf-8")
            pages = job / "pages.json"
            pages.write_text(json.dumps([{"page": 1, "text": "The study uses firm-level data and compares treated firms with a control group."}]), encoding="utf-8")
            base = {"workflowRoot": str(WORKFLOW), "jobDir": str(job), "extractPath": str(extract),
                    "pagesPath": str(pages), "name": "Reader", "date": "2026.10.19"}
            excel = bridge._render({**base, "stages": ["excel"],
                                    "pipelineStages": ["import", "prepare", "read", "excel", "lecture"]})
            book = openpyxl.load_workbook(excel["artifacts"][0]["path"], data_only=True)
            self.assertIn(book.active.cell(2, 16).value, (None, ""))
            self.assertEqual(book.active.cell(2, 17).value,
                             str(job / "artifacts" / "PPT_Author_2020.html"))
            self.assertFalse((job / "artifacts" / "PPT_Author_2020.html").exists())
            book.close()

    @unittest.skipUnless(WORKFLOW.is_dir(), "audited workflow directory is unavailable")
    def test_default_design_and_lecture_markdown_outputs_identify_segment_citations(self):
        record = sample_record()
        record["meta"]["citation_unit"] = "segment"
        record["research_questions"] += ["RQ2：政策通过什么机制影响创新？", "RQ3：政策效果是否因地区而异？"]
        record["hypotheses"] += ["H2（对应RQ2）：创新投入是可能的作用机制。",
                                "H3（对应RQ3）：政策效果在不同地区存在差异。"]
        record["economic_theories"] = ["创新扩散理论", "企业投资理论"]
        record["data"] += "样本覆盖多个地区和年份。变量口径依据论文说明构造。"
        notice = "来源为 Markdown；引用位置为 Markdown 段号，不是原 PDF 页码。"
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            job = root / "job"
            job.mkdir()
            extract = root / "extract.json"
            extract.write_text(json.dumps(record, ensure_ascii=False), encoding="utf-8")
            pages = job / "pages.json"
            pages.write_text(json.dumps([{
                "page": 1,
                "text": "The study uses firm-level data and compares treated firms with a control group.",
            }]), encoding="utf-8")
            base = {"workflowRoot": str(WORKFLOW), "jobDir": str(job),
                    "extractPath": str(extract), "pagesPath": str(pages)}

            for stage in ("design", "lecture"):
                with self.subTest(stage=stage):
                    result = bridge._render({**base, "stages": [stage]})
                    self.assertEqual(result["status"], "rendered", result)
                    artifact = result["artifacts"][0]
                    html = Path(artifact["path"]).read_text(encoding="utf-8")
                    self.assertIn(notice, html)
                    self.assertEqual(artifact["sha256"], bridge._sha256(html.encode("utf-8")))

        original_pdf_html = "<html><body>PDF source</body></html>"
        self.assertEqual(bridge._add_markdown_citation_notice(original_pdf_html, sample_record()),
                         original_pdf_html)

    @unittest.skipUnless(WORKFLOW.is_dir(), "audited workflow directory is unavailable")
    def test_excel_reuses_semantically_identical_file_and_rejects_conflict_or_tampering(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            job = root / "job"
            job.mkdir()
            extract = root / "extract.json"
            extract.write_text(json.dumps(sample_record(), ensure_ascii=False), encoding="utf-8")
            pages = job / "pages.json"
            pages.write_text(json.dumps([{"page": 1, "text": "The study uses firm-level data and compares treated firms with a control group."}]), encoding="utf-8")
            ctx = {"workflowRoot": str(WORKFLOW), "jobDir": str(job), "extractPath": str(extract),
                   "pagesPath": str(pages), "stages": ["excel"], "pipelineStages": ["import", "read", "excel"],
                   "name": "Reader", "date": "2026.10.19"}
            first = bridge._render(ctx)
            artifact = Path(first["artifacts"][0]["path"])
            first_hash = first["artifacts"][0]["sha256"]

            # Cross the XLSX core-property timestamp boundary. The bridge must
            # use its semantic fingerprint and verify the existing workbook,
            # not generate a new ZIP and report a false conflict.
            time.sleep(2.1)
            second = bridge._render(ctx)
            self.assertTrue(second["artifacts"][0]["reused"], second)
            self.assertEqual(second["artifacts"][0]["sha256"], first_hash)

            # Simulate a crash after the artifact write but before its receipt.
            (job / "render-receipt.json").unlink()
            recovered = bridge._render(ctx)
            self.assertTrue(recovered["artifacts"][0]["reused"], recovered)
            self.assertEqual(recovered["artifacts"][0]["sha256"], first_hash)

            changed = {**ctx, "name": "Different Reader"}
            with self.assertRaises(bridge.BridgeError) as caught:
                bridge._render(changed)
            self.assertEqual(caught.exception.status, "conflict")

            changed_record = sample_record()
            changed_record["claims"][0]["value"] += " changed but not displayed in Excel"
            changed_extract = root / "changed-extract.json"
            changed_extract.write_text(json.dumps(changed_record, ensure_ascii=False), encoding="utf-8")
            with self.assertRaises(bridge.BridgeError) as caught:
                bridge._render({**ctx, "extractPath": str(changed_extract)})
            self.assertEqual(caught.exception.status, "conflict")

            artifact.write_bytes(artifact.read_bytes() + b"tampered")
            with self.assertRaisesRegex(bridge.BridgeError, "hash与render receipt不一致"):
                bridge._render(ctx)

    @unittest.skipUnless(WORKFLOW.is_dir(), "audited workflow directory is unavailable")
    def test_publish_preview_is_stage_specific_and_offline(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            job = root / "job"
            job.mkdir()
            extract = root / "extract.json"
            extract.write_text(json.dumps(sample_record(), ensure_ascii=False), encoding="utf-8")
            pages = job / "pages.json"
            pages.write_text(json.dumps([{"page": 1, "text": "The study uses firm-level data and compares treated firms with a control group."}]), encoding="utf-8")
            base = {"workflowRoot": str(WORKFLOW), "jobDir": str(job), "extractPath": str(extract),
                    "pagesPath": str(pages), "name": "Reader", "date": "2026.10.19"}
            rendered = bridge._render({**base, "stages": ["design"]})
            self.assertEqual(rendered["status"], "rendered", rendered)

            blog = bridge._action_publish_preview({**base, "stages": ["blog"]})
            self.assertEqual(blog["status"], "preview_ready", blog)
            self.assertEqual(blog["stage"], "blog")
            self.assertEqual([f["targetName"] for f in blog["files"]], ["design_Author_2020.html"])
            self.assertEqual(blog["remoteNamePreflight"], "not_performed_no_network")
            self.assertNotIn("tsvPath", blog)
            self.assertNotIn("auth", json.dumps(blog))

            sheet = bridge._action_publish_preview({**base, "stages": ["sheet"]})
            self.assertEqual(sheet["status"], "preview_ready", sheet)
            self.assertEqual(sheet["stage"], "sheet")
            self.assertEqual(sheet["columnCount"], 21)
            self.assertEqual(len(Path(sheet["tsvPath"]).read_text(encoding="utf-8").rstrip("\n").split("\t")), 21)

    def test_publish_bundle_versions_when_blog_urls_change(self):
        with tempfile.TemporaryDirectory() as td:
            job = Path(td) / "job"
            cells_a = [f"c{i}" for i in range(21)]
            cells_b = list(cells_a)
            cells_b[17] = "https://example.invalid/new-design"
            reg_a = {"bridge": bridge.VERSION, "row": cells_a, "columnCount": 21}
            reg_b = {"bridge": bridge.VERSION, "row": cells_b, "columnCount": 21}
            tsv_a, registration_a, hash_a = bridge._persist_publish_bundle(job, cells_a, reg_a)
            tsv_b, registration_b, hash_b = bridge._persist_publish_bundle(job, cells_b, reg_b)
            self.assertNotEqual(hash_a, hash_b)
            self.assertNotEqual(tsv_a, tsv_b)
            self.assertNotEqual(registration_a, registration_b)
            self.assertTrue(tsv_a.is_file() and registration_a.is_file())
            self.assertTrue(tsv_b.is_file() and registration_b.is_file())
            self.assertIn("https://example.invalid/new-design", tsv_b.read_text(encoding="utf-8"))

    def test_sheet_write_requires_explicit_server_append_semantics(self):
        append_tool = {
            "description": "Append rows to the end of the selected sheet; existing rows are preserved.",
            "inputSchema": {"properties": {"file_id": {}, "sheet_name": {}, "values": {}, "range": {}},
                            "required": ["file_id", "sheet_name", "values"]},
        }
        range_tool = {
            "description": "Update cells in the requested spreadsheet range.",
            "inputSchema": {"properties": {"file_id": {}, "sheet_name": {}, "values": {}, "range": {}},
                            "required": ["file_id", "sheet_name", "values", "range"]},
        }
        vague_tool = {
            "description": "This tool can append values or update a selected range.",
            "inputSchema": {"properties": {"file_id": {}, "sheet_name": {}, "values": {}, "range": {}},
                            "required": ["file_id", "sheet_name", "values"]},
        }
        self.assertEqual(bridge._sheet_write_strategy(append_tool), "append")
        self.assertIsNone(bridge._sheet_write_strategy(range_tool))
        self.assertIsNone(bridge._sheet_write_strategy(vague_tool))
        args = bridge._sheet_tool_args(append_tool, {
            "file_id": "doc", "sheet_name": "Sheet", "values": [["x"]],
        })
        self.assertNotIn("range", args)

    def test_non_append_sheet_tool_is_never_called(self):
        class FakeClient:
            writes = 0
            calls = []

            def __init__(self, endpoint, token):
                self.endpoint = endpoint
                self.token = token

            def initialize(self):
                return None

            def list_tools(self):
                return [
                    {"name": "search_space_file", "inputSchema": {
                        "properties": {"keyword": {}}, "required": ["keyword"]}},
                    {"name": "get_content", "inputSchema": {
                        "properties": {"file_id": {}, "sheet_name": {}},
                        "required": ["file_id", "sheet_name"]}},
                    {"name": "batch_update_sheet_range",
                     "description": "Update cells in the requested spreadsheet range.",
                     "inputSchema": {"properties": {"file_id": {}, "sheet_name": {},
                                                       "values": {}, "range": {}},
                                     "required": ["file_id", "sheet_name", "values", "range"]}},
                ]

            def call_tool(self, name, args):
                self.calls.append(name)
                if name == "search_space_file":
                    return {"node_id": "mock-doc", "url": "https://docs.qq.com/sheet/DTHRhWHFwaGJPTWtL"}
                if name == "get_content":
                    return {"values": []}
                if name == "batch_update_sheet_range":
                    type(self).writes += 1
                    raise AssertionError("unsafe write must not be called")
                raise AssertionError(f"unexpected mock tool: {name}")

        fake_log = SimpleNamespace(MCP_ENDPOINT="https://example.invalid/mcp",
                                   DOC_TITLE="mock", McpClient=FakeClient)
        with tempfile.TemporaryDirectory() as td:
            root = Path(td) / "workflow"
            root.mkdir()
            job = Path(td) / "job"
            job.mkdir()
            receipt_path = job / "receipt.json"
            ctx = {"workflowRoot": str(root), "auth": {"tencentToken": "mock-token"},
                   "sheetName": "Sheet"}
            receipt = {"bridge": bridge.VERSION, "entries": {}}
            with patch.object(bridge, "_import_file", return_value=fake_log):
                result = bridge._publish_sheet(ctx, [f"c{i}" for i in range(21)],
                                               job, receipt, receipt_path)
            self.assertEqual(result["status"], "needs_setup", result)
            self.assertEqual(FakeClient.writes, 0)
            self.assertNotIn("batch_update_sheet_range", FakeClient.calls)
            self.assertFalse(receipt_path.exists())

    def test_sheet_client_preserves_the_verified_authorization_header_scheme(self):
        seen = []
        class FakeClient:
            def __init__(self, endpoint, token):
                self.token = token
            def _headers(self):
                return {"Authorization": "Bearer " + self.token}
            def initialize(self):
                seen.append(self._headers()["Authorization"])
            def list_tools(self):
                return []
        fake_log = SimpleNamespace(MCP_ENDPOINT="https://example.invalid/mcp", DOC_TITLE="mock", McpClient=FakeClient)
        with tempfile.TemporaryDirectory() as td:
            root = Path(td) / "workflow"
            root.mkdir()
            job = Path(td) / "job"
            job.mkdir()
            for scheme in ["raw", "bearer"]:
                ctx = {"workflowRoot": str(root), "auth": {"tencentToken": "fixture", "tencentAuthScheme": scheme}}
                with patch.object(bridge, "_import_file", return_value=fake_log):
                    result = bridge._publish_sheet(ctx, ["x"] * 21, job, {"entries": {}}, job / "receipt.json")
                self.assertEqual(result["status"], "needs_setup")
            self.assertEqual(seen, ["fixture", "Bearer fixture"])
            self.assertFalse((job / "receipt.json").exists())

    def test_sheet_target_must_match_configured_url_uniquely(self):
        target = "https://docs.qq.com/sheet/DTHRhWHFwaGJPTWtL"
        find = bridge._find_bound_document_id
        wrong = {"node_id": "wrong", "url": "https://docs.qq.com/sheet/other"}
        right = {"node_id": "right", "url": target + "?tab=one"}
        self.assertEqual(find([wrong, right], target), "right")
        self.assertIsNone(find([wrong, {"node_id": "unknown"}], target))
        self.assertIsNone(find([right, {"file_id": "different", "url": target}], target))
        self.assertIsNone(find({"node_id": "parent", "child": {"url": target}}, target))
        self.assertIsNone(find({"node_id": "evil", "url": target.replace("docs.qq.com", "docs.qq.com.evil.invalid")}, target))

    def test_missing_blog_credentials_is_needs_auth_without_request(self):
        with tempfile.TemporaryDirectory() as td:
            result = bridge._publish_blog(
                {"workflowRoot": str(WORKFLOW), "auth": {"blogAccount": "user"}, "stages": ["blog"]},
                sample_record(), Path(td), {"bridge": bridge.VERSION, "entries": {}}, Path(td) / "receipt.json",
            )
            self.assertEqual(result["status"], "needs_auth")
            self.assertNotIn("user", json.dumps(result))

    def test_redaction_removes_stdin_credentials(self):
        ctx = {"auth": {"blogPassword": "secret-blog-value", "tencentToken": "secret-token-value"}}
        result = bridge._redact_result({"message": "failed secret-blog-value secret-token-value"}, ctx)
        self.assertNotIn("secret-blog-value", json.dumps(result))
        self.assertNotIn("secret-token-value", json.dumps(result))


if __name__ == "__main__":
    unittest.main(verbosity=2)
