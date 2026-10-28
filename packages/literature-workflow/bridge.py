#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""DeepSeek Harness JSON bridge for the read-only-audited literature workflow.

Invocation: ``python -X utf8 bridge.py <action>`` with one JSON object on stdin.
The process emits exactly one JSON object on stdout. Credentials, when needed,
belong only in stdin ``auth`` and are never copied into a receipt or response.
"""
from __future__ import annotations

import contextlib
import copy
import datetime as _dt
import hashlib
import importlib.util
import io
import json
import os
import re
import subprocess
import sys
import tempfile
import time
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urlsplit

VERSION = "0.1.0"
LOCAL_STAGES = {"excel", "design", "lecture"}
PUBLISH_STAGES = {"blog", "sheet"}
ALL_STAGES = LOCAL_STAGES | PUBLISH_STAGES
EXPECTED_COLUMNS = 21
DEFAULT_BLOG_FOLDER = "master_fin"
PLACEHOLDER_RE = re.compile(
    r"(?:\.\.\.|…|10\.x{3,}|\bxxx+\b|\bJournal Name\b|"
    r"Author,\s*A\.\s*A\.|论文标题（原文）|200-300\s*字中文摘要|"
    r"RQ\d+\s*:\s*\.{2,}|H\d+[^\n]{0,24}\.\.\.)",
    re.IGNORECASE,
)
NUMBER_RE = re.compile(r"(?<![A-Za-z])[-+]?\d+(?:\.\d+)?\s*%?")
DOI_RE = re.compile(r"\b10\.\d{4,9}/[-._;()/:A-Z0-9]+", re.IGNORECASE)
YEAR_RE = re.compile(r"\b(?:19|20)\d{2}\b")
_SECRET_KEYS = {"blogpassword", "password", "tencenttoken", "token", "access_token", "access-token"}


class BridgeError(RuntimeError):
    def __init__(self, message: str, status: str = "error") -> None:
        super().__init__(message)
        self.status = status


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _add_markdown_citation_notice(html: str, record: dict[str, Any]) -> str:
    """Label segment citations in default decks when the source is Markdown."""
    meta = record.get("meta") if isinstance(record, dict) else None
    if not isinstance(meta, dict) or meta.get("citation_unit") != "segment":
        return html
    notice_text = "来源为 Markdown；引用位置为 Markdown 段号，不是原 PDF 页码。"
    if notice_text in html:
        return html
    notice = (
        '<aside class="source-citation-notice" role="note" '
        'style="padding:.55rem 1rem;background:#fff2cc;color:#553c00;'
        'border-bottom:1px solid #d3b35a;font:600 14px/1.5 sans-serif;'
        'text-align:center">'
        f"{notice_text}</aside>"
    )
    body_open = re.search(r"<body\b[^>]*>", html, flags=re.IGNORECASE)
    if body_open:
        return html[:body_open.end()] + notice + html[body_open.end():]
    body_close = re.search(r"</body\s*>", html, flags=re.IGNORECASE)
    if body_close:
        return html[:body_close.start()] + notice + html[body_close.start():]
    return notice + html


def _json_bytes(value: Any) -> bytes:
    return (json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n").encode("utf-8")


def _path(value: Any, label: str, *, must_exist: bool = False) -> Path:
    if not isinstance(value, str) or not value.strip():
        raise BridgeError(f"缺少 {label} 路径。", "needs_setup")
    p = Path(value).expanduser().resolve()
    if must_exist and not p.exists():
        raise BridgeError(f"{label} 不存在。", "needs_setup")
    return p


def _within(child: Path, parent: Path) -> bool:
    try:
        child.resolve().relative_to(parent.resolve())
        return True
    except ValueError:
        return False


def _job_dir(ctx: dict[str, Any], root: Path) -> Path:
    job = _path(ctx.get("jobDir"), "jobDir")
    if _within(job, root):
        rel = job.relative_to(root.resolve())
        if not rel.parts or rel.parts[0].casefold() != ".plugin-state":
            raise BridgeError("jobDir 位于工作流源码目录内；仅允许写入工作流的 .plugin-state 子目录。", "needs_setup")
    return job


def _source_path(root: Path, rel: str) -> Path:
    p = (root / rel).resolve()
    if not _within(p, root):
        raise BridgeError("工作流源码路径越界。", "needs_setup")
    return p


def _persist_immutable(entries: Iterable[tuple[Path, bytes]]) -> dict[str, bool]:
    """Create outputs once; an existing identical file is safely reused."""
    planned = list(entries)
    for path, data in planned:
        if path.exists() and path.read_bytes() != data:
            raise BridgeError(f"输出已存在且内容不同，拒绝覆盖：{path}", "conflict")
    result: dict[str, bool] = {}
    for path, data in planned:
        path.parent.mkdir(parents=True, exist_ok=True)
        if path.exists():
            result[str(path)] = True
            continue
        try:
            with path.open("xb") as f:
                f.write(data)
                f.flush()
                os.fsync(f.fileno())
            result[str(path)] = False
        except FileExistsError:
            if path.read_bytes() != data:
                raise BridgeError(f"输出并发冲突，拒绝覆盖：{path}", "conflict")
            result[str(path)] = True
    return result


def _atomic_json(path: Path, value: Any, *, owned_marker: str | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        try:
            old = json.loads(path.read_text(encoding="utf-8"))
        except Exception as exc:
            raise BridgeError(f"已有回执不可解析，拒绝覆盖：{path}", "conflict") from exc
        if owned_marker and old.get("bridge") != owned_marker:
            raise BridgeError(f"目标文件不是本桥接器回执，拒绝覆盖：{path}", "conflict")
    fd, temp_name = tempfile.mkstemp(prefix=".bridge-", suffix=".tmp", dir=str(path.parent))
    try:
        with os.fdopen(fd, "wb") as f:
            f.write(_json_bytes(value))
            f.flush()
            os.fsync(f.fileno())
        os.replace(temp_name, path)
    finally:
        try:
            os.unlink(temp_name)
        except FileNotFoundError:
            pass


def _load_json(path: Path, label: str) -> dict[str, Any]:
    try:
        obj = json.loads(path.read_text(encoding="utf-8-sig"))
    except Exception as exc:
        raise BridgeError(f"{label} 不是可读取的 JSON：{type(exc).__name__}", "invalid_input") from None
    if not isinstance(obj, dict):
        raise BridgeError(f"{label} 顶层必须是 JSON 对象。", "invalid_input")
    return obj


def _normalise_extract(obj: dict[str, Any]) -> dict[str, Any]:
    """Accept the workflow's flat schema and {meta, record} wrapper."""
    if isinstance(obj.get("record"), dict):
        record = copy.deepcopy(obj["record"])
        if "meta" not in record and isinstance(obj.get("meta"), dict):
            record["meta"] = copy.deepcopy(obj["meta"])
        if "evidence" not in record and isinstance(obj.get("evidence"), list):
            record["evidence"] = copy.deepcopy(obj["evidence"])
        if "claims" not in record and isinstance(obj.get("claims"), list):
            record["claims"] = copy.deepcopy(obj["claims"])
    else:
        record = copy.deepcopy(obj)
    meta = record.get("meta") if isinstance(record.get("meta"), dict) else {}
    if not record.get("year") and meta.get("year"):
        record["year"] = str(meta["year"])
    if not record.get("authors") and meta.get("authors_full"):
        record["authors"] = str(meta["authors_full"])
    return record


def _load_record(ctx: dict[str, Any]) -> tuple[dict[str, Any], Path]:
    root = _path(ctx.get("workflowRoot"), "workflowRoot", must_exist=True)
    extract_path = _path(ctx.get("extractPath"), "extractPath", must_exist=True)
    obj = _load_json(extract_path, "extractPath")
    return _normalise_extract(obj), root


def _import_file(path: Path, name: str):
    if not path.is_file():
        raise BridgeError(f"工作流缺少所需模块：{path}", "needs_setup")
    spec = importlib.util.spec_from_file_location(name, str(path))
    if spec is None or spec.loader is None:
        raise BridgeError(f"无法载入工作流模块：{path}", "needs_setup")
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


def _load_page_texts(job: Path, explicit_dir: str | None = None) -> dict[int, str]:
    folder = Path(explicit_dir).resolve() if explicit_dir else job / "pages"
    if folder.is_file():
        try:
            payload = json.loads(folder.read_text(encoding="utf-8-sig"))
        except Exception as exc:
            raise BridgeError(f"pagesPath 无法读取：{type(exc).__name__}。", "invalid_input") from None
        if not isinstance(payload, list):
            raise BridgeError("pagesPath JSON 必须是 [{page,text}, ...] 数组。", "invalid_input")
        result: dict[int, str] = {}
        for item in payload:
            if isinstance(item, dict) and isinstance(item.get("page"), int) and isinstance(item.get("text"), str):
                result[item["page"]] = item["text"]
        return result
    if not folder.is_dir():
        return {}
    pages: dict[int, str] = {}
    for p in sorted(folder.glob("page-*.txt")):
        m = re.fullmatch(r"page-(\d+)\.txt", p.name, flags=re.IGNORECASE)
        if not m:
            continue
        body = p.read_text(encoding="utf-8", errors="replace")
        clean = []
        for line in body.splitlines():
            clean.append(re.sub(r"^\s*\d{1,6}\s*[:\t]\s?", "", line))
        pages[int(m.group(1))] = "\n".join(clean)
    return pages


def _norm_quote(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip().casefold()


def _is_pending(value: Any) -> bool:
    s = str(value or "").strip().casefold()
    return any(x in s for x in ("待核对", "待补", "未报告", "not reported", "n/a"))


def _meaningful(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return bool(value.strip())
    if isinstance(value, (list, tuple, dict)):
        return bool(value)
    return True


def _walk_strings(value: Any, prefix: str = "") -> Iterable[tuple[str, str]]:
    if isinstance(value, dict):
        for key, child in value.items():
            if key in {"_meta", "_validation", "_comment"}:
                continue
            path = f"{prefix}.{key}" if prefix else str(key)
            yield from _walk_strings(child, path)
    elif isinstance(value, (list, tuple)):
        for index, child in enumerate(value):
            yield from _walk_strings(child, f"{prefix}.{index}")
    elif isinstance(value, str):
        yield prefix, value


def _field_evidence(field: str, global_evidence: list[dict], local_evidence: list[dict] | None = None) -> list[dict]:
    refs = list(global_evidence)
    if isinstance(local_evidence, list):
        refs.extend(x for x in local_evidence if isinstance(x, dict))
    return [e for e in refs if e.get("field") == field or field.startswith(str(e.get("field", "")) + ".")]


def _has_path(record: dict[str, Any], field: str) -> bool:
    cur: Any = record
    for part in field.split("."):
        if isinstance(cur, dict) and part in cur:
            cur = cur[part]
        elif isinstance(cur, list) and part.isdigit() and int(part) < len(cur):
            cur = cur[int(part)]
        else:
            return False
    return True


def validate_record(record: dict[str, Any], page_texts: dict[int, str],
                    workflow_root: Path | None = None) -> dict[str, Any]:
    """Validate legacy worksheet fields plus evidence/page references."""
    errors: list[str] = []
    warnings: list[str] = []
    required = (
        "apa_citation", "authors", "year", "abstract", "research_questions", "hypotheses",
        "core_theory", "contribution", "data", "limitations", "future_research", "fun_story",
        "confusions", "why_top", "core_variables", "economic_concepts", "economic_theories",
    )
    for key in required:
        if not _meaningful(record.get(key)):
            errors.append(f"缺少21列抽取字段：{key}")

    meta = record.get("meta")
    if not isinstance(meta, dict):
        errors.append("meta 必须是对象，至少包含 title、authors_full、year。")
        meta = {}
    for key in ("title", "authors_full", "year"):
        if not _meaningful(meta.get(key)):
            errors.append(f"meta.{key} 缺失。")
    if not isinstance(record.get("design"), dict):
        errors.append("design 必须是对象；缺项可写明【待补】，不能省略整个结构。")
    else:
        design = record["design"]
        for key in ("variables", "mechanism", "heterogeneity", "robustness", "key_coefficients"):
            if key in design and not isinstance(design[key], list):
                errors.append(f"design.{key} 必须是数组。")
        ident = design.get("identification")
        if ident is not None and not isinstance(ident, dict):
            errors.append("design.identification 必须是对象。")

    # Reuse the legacy column validator when its local dependency is available.
    try:
        root = Path(workflow_root).resolve() if workflow_root else Path(__file__).resolve().parents[2]
        excel = _import_file(root / "src" / "excel" / "make_excel.py", "_dsh_lit_make_excel")
        errors.extend(excel.validate(record))
    except Exception as exc:
        warnings.append(f"旧工作流 make_excel.validate 未载入：{type(exc).__name__}")

    for field, text in _walk_strings(record):
        # Explicit pending markers are data-status labels, not sample prose.
        sample_text = text.replace("【待补】", "").replace("【待核对】", "")
        if PLACEHOLDER_RE.search(sample_text):
            errors.append(f"字段 {field} 看起来仍含 schema 示例占位文本。")

    evidence_raw = record.get("evidence", [])
    if evidence_raw is None:
        evidence_raw = []
    if not isinstance(evidence_raw, list):
        errors.append("evidence 必须是 [{field,page,quote}, ...] 数组。")
        evidence_raw = []
    evidence: list[dict] = [e for e in evidence_raw if isinstance(e, dict)]
    evidence_unverified = not page_texts
    if evidence_unverified:
        warnings.append("未提供 pagesPath；页码与quote结构已读取，但原文证据无法在本次验证。")
    for i, item in enumerate(evidence_raw):
        if not isinstance(item, dict):
            errors.append(f"evidence[{i}] 必須是对象。")
            continue
        field, page, quote = item.get("field"), item.get("page"), item.get("quote")
        if not isinstance(field, str) or not field.strip():
            errors.append(f"evidence[{i}].field 缺失。")
            continue
        if not _has_path(record, field):
            errors.append(f"evidence[{i}].field 不对应 extract 中的字段：{field}。")
            continue
        if not isinstance(page, int) or isinstance(page, bool) or page < 1:
            errors.append(f"evidence[{i}].page 必须是从1开始的整数。")
            continue
        if not isinstance(quote, str) or not quote.strip():
            errors.append(f"evidence[{i}].quote 缺失。")
            continue
        source = page_texts.get(page)
        if source is None and page_texts:
            errors.append(f"evidence[{i}] 引用了未准备的第 {page} 页。")
        elif source is not None and _norm_quote(quote) not in _norm_quote(source):
            errors.append(f"evidence[{i}] 的 quote 未在第 {page} 页文本中找到。")

    claims = record.get("claims")
    if claims is not None:
        if not isinstance(claims, list):
            errors.append("claims 必须是数组。")
        else:
            for i, claim in enumerate(claims):
                if not isinstance(claim, dict):
                    errors.append(f"claims[{i}] 必须是对象。")
                    continue
                field = claim.get("field")
                val = claim.get("value", claim.get("claim", claim.get("text")))
                local = claim.get("evidence")
                if not isinstance(field, str) or not field.strip() or not _meaningful(val):
                    errors.append(f"claims[{i}] 需提供 field 与 value/claim/text。")
                    continue
                refs = _field_evidence(field, evidence, local if isinstance(local, list) else None)
                if not refs:
                    errors.append(f"主要论断 claims[{i}] ({field}) 缺少 evidence；数字结果请标【待核对】，不得补造。")
                if not _has_path(record, field):
                    errors.append(f"claims[{i}].field 不对应 extract 中的字段：{field}。")
                if isinstance(local, list):
                    for j, ref in enumerate(local):
                        if not isinstance(ref, dict):
                            errors.append(f"claims[{i}].evidence[{j}] 必须是对象。")
                            continue
                        ref_field, page, quote = ref.get("field"), ref.get("page"), ref.get("quote")
                        if ref_field != field:
                            errors.append(f"claims[{i}].evidence[{j}].field 应与 claim field 一致。")
                        if not isinstance(page, int) or isinstance(page, bool) or page < 1:
                            errors.append(f"claims[{i}].evidence[{j}].page 必须是从1开始的整数。")
                        elif page_texts and page not in page_texts:
                            errors.append(f"claims[{i}].evidence[{j}] 引用了未准备的第 {page} 页。")
                        if not isinstance(quote, str) or not quote.strip():
                            errors.append(f"claims[{i}].evidence[{j}].quote 缺失。")
                        elif page in page_texts and _norm_quote(quote) not in _norm_quote(page_texts[page]):
                            errors.append(f"claims[{i}] 的 evidence quote 未在第 {page} 页找到。")
                    evidence.extend(ref for ref in local if isinstance(ref, dict))
                for ref in refs:
                    if ref.get("field") != field and not field.startswith(str(ref.get("field", "")) + "."):
                        continue
                    page = ref.get("page")
                    quote = ref.get("quote")
                    if not isinstance(page, int) or page not in page_texts or not isinstance(quote, str):
                        continue  # detailed reference errors are emitted above for global evidence
                    if _norm_quote(quote) not in _norm_quote(page_texts[page]):
                        errors.append(f"claims[{i}] 的 evidence quote 未在第 {page} 页找到。")
    else:
        warnings.append("未提供 claims 清单；无法逐条核验所有文字性论断。")

    # A numerical coefficient/result is never silently accepted without a matching page quote.
    design = record.get("design") if isinstance(record.get("design"), dict) else {}
    coefs = design.get("key_coefficients") if isinstance(design.get("key_coefficients"), list) else []
    for idx, coef in enumerate(coefs):
        if not isinstance(coef, dict):
            continue
        for key in ("coef", "se", "stars", "result"):
            value = coef.get(key)
            if not _meaningful(value) or _is_pending(value) or not NUMBER_RE.search(str(value)):
                if _is_pending(value):
                    warnings.append(f"design.key_coefficients.{idx}.{key} 标记待核对，未当作已证实结果。")
                continue
            field = f"design.key_coefficients.{idx}.{key}"
            if not _field_evidence(field, evidence):
                errors.append(f"数值结果 {field} 缺少页码证据；未核实前请标【待核对】，不要编造。")

    return {"ok": not errors, "evidenceStatus": "unavailable" if evidence_unverified else "verified",
            "errors": errors, "warnings": warnings,
            "checkedEvidence": len(evidence), "availablePages": sorted(page_texts)}


def _metadata(pdf_path: Path, converter: Any, pymupdf: Any, page_texts: list[str]) -> dict[str, Any]:
    first = page_texts[0] if page_texts else ""
    title = None
    title_source = None
    try:
        with pymupdf.open(str(pdf_path)) as doc:
            candidate = converter._first_page_title(doc)
            if isinstance(candidate, str) and candidate.strip():
                title, title_source = candidate.strip(), "PDF first-page layout extraction"
    except Exception:
        pass
    dois = []
    for match in DOI_RE.findall(first):
        d = match.rstrip(".,;:)】]}")
        if d.casefold() not in {x.casefold() for x in dois}:
            dois.append(d)
    years = sorted(set(YEAR_RE.findall(first)))
    return {
        "title": title,
        "titleSource": title_source,
        "doi": dois[0] if len(dois) == 1 else None,
        "doiCandidates": dois,
        "year": int(years[0]) if len(years) == 1 else None,
        "yearCandidates": [int(x) for x in years],
    }


def _action_probe(ctx: dict[str, Any]) -> dict[str, Any]:
    root = _path(ctx.get("workflowRoot"), "workflowRoot", must_exist=True)
    schema = _source_path(root, "templates/extract_schema.json")
    py_mods = {}
    for name in ("pymupdf", "pymupdf4llm", "openpyxl", "requests"):
        try:
            py_mods[name] = bool(importlib.util.find_spec(name))
        except Exception:
            py_mods[name] = False
    return {
        "ok": schema.is_file(), "status": "ready" if schema.is_file() else "needs_setup",
        "workflowRoot": str(root), "schemaPath": str(schema),
        "modules": py_mods,
        "capabilities": {"prepare": "PDF text extraction; no Zotero write",
                         "validate": "legacy 21-column fields plus page evidence",
                         "render": sorted(LOCAL_STAGES), "publish": sorted(PUBLISH_STAGES)},
        "limitations": ["No paper-search/download integration is present in the audited workflow.",
                        "Zotero lookup/write is intentionally outside this bridge."]
    }


def _action_prepare(ctx: dict[str, Any]) -> dict[str, Any]:
    root = _path(ctx.get("workflowRoot"), "workflowRoot", must_exist=True)
    pdf = _path(ctx.get("pdfPath"), "pdfPath", must_exist=True)
    if pdf.suffix.lower() != ".pdf":
        raise BridgeError("pdfPath 必须指向 PDF 文件。", "invalid_input")
    job = _job_dir(ctx, root)
    converter_path = _source_path(root, "src/pdf2md/convert.py")
    pdf2md_dir = str(converter_path.parent)
    if pdf2md_dir not in sys.path:
        sys.path.insert(0, pdf2md_dir)
    converter = _import_file(converter_path, "_dsh_lit_pdf_convert")
    try:
        import pymupdf
    except ImportError:
        try:
            import fitz as pymupdf
        except ImportError:
            raise BridgeError("缺少 PyMuPDF；无法准备逐页文本。", "needs_setup") from None

    noise: list[str] = []
    engine = ctx.get("engine", "auto")
    if engine not in {"auto", "pymupdf4llm", "blocks"}:
        raise BridgeError("engine 仅支持 auto / pymupdf4llm / blocks。", "invalid_input")
    try:
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            converted = converter.convert_pdf(str(pdf), engine=engine, quiet=True, library_noise=noise)
        with pymupdf.open(str(pdf)) as doc:
            page_texts = [page.get_text("text", sort=True) for page in doc]
            page_count = doc.page_count
    except Exception as exc:
        raise BridgeError(f"PDF 准备失败：{type(exc).__name__}。", "conversion_failed") from None
    if not isinstance(converted, dict) or not isinstance(converted.get("markdown"), str):
        raise BridgeError("转换器未返回完整 Markdown。", "conversion_failed")
    markdown = converted["markdown"]
    if not markdown.strip():
        raise BridgeError("转换结果为空，拒绝标记准备成功。", "conversion_failed")
    pages_dir = job / "pages"
    pages_json = job / "pages.json"
    text_entries: list[tuple[Path, bytes]] = []
    for idx, text in enumerate(page_texts, start=1):
        numbered = "\n".join(f"{line_no:04d}: {line}" for line_no, line in enumerate(text.splitlines(), start=1))
        if text.endswith("\n"):
            numbered += "\n"
        text_entries.append((pages_dir / f"page-{idx:04d}.txt", numbered.encode("utf-8")))
    scan_pages = [i for i, text in enumerate(page_texts, start=1)
                  if len(text.strip()) < getattr(converter, "MIN_PAGE_RAW", 24)]
    metadata = _metadata(pdf, converter, pymupdf, page_texts)
    md_path = job / "full.md"
    meta_path = job / "metadata.json"
    meta_result = {
        "pdfPath": str(pdf), "markdownPath": str(md_path), "pagesPath": str(pages_json),
        "pageTextDir": str(pages_dir), "extractSchemaPath": str(_source_path(root, "templates/extract_schema.json")),
        "pageCount": page_count, "scanPages": scan_pages,
        "title": metadata["title"], "titleSource": metadata["titleSource"],
        "doi": metadata["doi"], "doiCandidates": metadata["doiCandidates"],
        "year": metadata["year"], "yearCandidates": metadata["yearCandidates"],
        "converter": converted.get("engine"), "fallbackPages": converted.get("fallback_pages", []),
        "droppedPages": converted.get("dropped_pages", []),
    }
    page_payload = [{"page": i, "text": text} for i, text in enumerate(page_texts, start=1)]
    _persist_immutable([(md_path, markdown.encode("utf-8")), *text_entries,
                        (pages_json, _json_bytes(page_payload)), (meta_path, _json_bytes(meta_result))])
    return {"ok": True, "status": "ready", **meta_result}


def _action_validate(ctx: dict[str, Any]) -> dict[str, Any]:
    record, root = _load_record(ctx)
    job = _job_dir(ctx, root)
    page_texts = _load_page_texts(job, ctx.get("pagesPath") or ctx.get("pageTextDir"))
    result = validate_record(record, page_texts, root)
    return {"status": ("invalid" if not result["ok"] else ("valid" if result["evidenceStatus"] == "verified" else "valid_needs_evidence")),
            "extractPath": str(_path(ctx.get("extractPath"), "extractPath")),
            "schemaPath": str(_source_path(root, "templates/extract_schema.json")), **result}


def _stage_list(ctx: dict[str, Any]) -> list[str]:
    raw = ctx.get("stages")
    if raw is None:
        raw = ctx.get("selectedStages")
    if raw is None:
        return []
    if not isinstance(raw, list) or not all(isinstance(x, str) for x in raw):
        raise BridgeError("stages/selectedStages 必须是字符串数组。", "invalid_input")
    stages = list(dict.fromkeys(x.strip().lower() for x in raw if x.strip()))
    unknown = sorted(set(stages) - ALL_STAGES)
    if unknown:
        raise BridgeError("包含不支持的 stage：" + ", ".join(unknown), "invalid_input")
    return stages


def _safe_author_year(record: dict[str, Any]) -> tuple[str, str]:
    meta = record.get("meta") if isinstance(record.get("meta"), dict) else {}
    surname = str(meta.get("first_author_surname") or meta.get("first_author_surname_en") or "Paper")
    surname = re.sub(r"[^A-Za-z0-9'-]+", "_", surname).strip("_-")[:60] or "Paper"
    year_raw = str(record.get("year") or meta.get("year") or "")
    year_match = re.search(r"\b(?:19|20)\d{2}\b", year_raw)
    return surname, year_match.group(0) if year_match else "0000"


def _load_render_receipt(job: Path) -> tuple[dict[str, Any], Path]:
    path = job / "render-receipt.json"
    if not path.exists():
        return {"bridge": VERSION, "rendered": {}}, path
    receipt = _load_json(path, "render receipt")
    if receipt.get("bridge") != VERSION or not isinstance(receipt.get("rendered"), dict):
        raise BridgeError("已有render receipt不属于本桥接器或格式不可用。", "conflict")
    return receipt, path


def _save_render_entry(job: Path, stage: str, entry: dict[str, Any]) -> None:
    # Re-read before merging so a completed card does not erase other stage
    # receipts written since this bridge invocation began.
    receipt, path = _load_render_receipt(job)
    receipt["rendered"][stage] = entry
    _atomic_json(path, receipt, owned_marker=VERSION)


def _verified_rendered_path(receipt: dict[str, Any], artifacts_dir: Path, stage: str) -> Path | None:
    info = (receipt.get("rendered") or {}).get(stage)
    if not isinstance(info, dict) or not info.get("path"):
        return None
    if not isinstance(info["path"], str):
        raise BridgeError(f"render receipt中的{stage}路径格式无效。", "conflict")
    path = Path(info["path"]).resolve()
    if not _within(path, artifacts_dir.resolve()) or not path.is_file():
        raise BridgeError(f"render receipt中的{stage}路径无效或文件缺失。", "conflict")
    if not info.get("sha256") or _sha256(path.read_bytes()) != info.get("sha256"):
        raise BridgeError(f"render receipt中的{stage}文件hash不匹配。", "conflict")
    return path


def _excel_deck_links(record: dict[str, Any], job: Path, receipt: dict[str, Any],
                      pipeline_stages: list[str] | None) -> dict[str, str]:
    artifacts_dir = job / "artifacts"
    surname, year = _safe_author_year(record)
    future_names = {"design": f"design_{surname}_{year}.html", "lecture": f"PPT_{surname}_{year}.html"}
    stage_indexes: dict[str, int] = {}
    if pipeline_stages is not None:
        for index, stage in enumerate(pipeline_stages):
            stage_indexes.setdefault(stage, index)
    excel_index = stage_indexes.get("excel")
    links: dict[str, str] = {}
    for stage in ("design", "lecture"):
        existing = _verified_rendered_path(receipt, artifacts_dir, stage)
        if existing is not None:
            links[stage] = str(existing)
            continue
        stage_index = stage_indexes.get(stage)
        if excel_index is None or stage_index is None or stage_index <= excel_index:
            links[stage] = ""
            continue
        future_path = artifacts_dir / future_names[stage]
        if future_path.exists():
            # A future-stage filename without a verified receipt may be stale
            # output from an interrupted or older render. Never link it as if
            # the selected later card were guaranteed to create this content.
            raise BridgeError(f"计划中的{stage}文件已存在但没有有效render receipt：{future_path.name}", "conflict")
        links[stage] = str(future_path)
    return links


def _excel_semantics_match(path: Path, excel: Any, record: dict[str, Any]) -> bool:
    """Ignore XLSX ZIP timestamps; reuse only when every visible cell matches."""
    workbook = None
    try:
        import openpyxl
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        columns = excel.COLUMNS
        if sheet.max_column != EXPECTED_COLUMNS or sheet.max_row != 2 or len(columns) != EXPECTED_COLUMNS:
            return False
        for column, (key, header, kind, *_rest) in enumerate(columns, start=1):
            if sheet.cell(1, column).value != header:
                return False
            expected = excel.render_cell(record.get(key), kind)
            actual = sheet.cell(2, column).value
            if ("" if actual is None else str(actual)) != ("" if expected is None else str(expected)):
                return False
        return True
    except Exception:
        return False
    finally:
        if workbook is not None:
            workbook.close()


def _pipeline_stages(ctx: dict[str, Any]) -> list[str] | None:
    raw = ctx.get("pipelineStages")
    if raw is None:
        return None
    if not isinstance(raw, list) or not all(isinstance(stage, str) for stage in raw):
        raise BridgeError("pipelineStages 必须是字符串数组。", "invalid_input")
    return [stage.strip().lower() for stage in raw]


def _render(ctx: dict[str, Any]) -> dict[str, Any]:
    record, root = _load_record(ctx)
    job = _job_dir(ctx, root)
    stages = [s for s in _stage_list(ctx) if s in LOCAL_STAGES]
    if not stages:
        return {"ok": True, "status": "skipped", "selectedStages": [], "artifacts": []}
    if len(stages) != 1:
        raise BridgeError("render 每次只接受一个 stage：excel、design 或 lecture。", "invalid_input")
    validation = validate_record(record, _load_page_texts(job, ctx.get("pagesPath") or ctx.get("pageTextDir")), root)
    if not validation["ok"]:
        return {"ok": False, "status": "invalid", "validation": validation, "artifacts": []}
    render_receipt, render_receipt_path = _load_render_receipt(job)
    pipeline_stages = _pipeline_stages(ctx)
    artifacts_dir = job / "artifacts"
    artifacts_dir.mkdir(parents=True, exist_ok=True)
    surname, year = _safe_author_year(record)
    md_path = job / "full.md"
    if not md_path.is_file() and ctx.get("markdownPath"):
        md_path = _path(ctx["markdownPath"], "markdownPath", must_exist=True)
    produced: list[dict[str, Any]] = []
    logs = io.StringIO()
    custom_plan_path = ctx.get('stagePlanPath')
    if custom_plan_path:
        stage = stages[0]
        if stage not in ('design','lecture'):
            raise BridgeError('此阶段不支持自定义演示编排。', 'invalid_input')
        plan_path = _path(custom_plan_path, 'stagePlanPath', must_exist=True)
        if not _within(plan_path, job):
            raise BridgeError('自定义编排文件必须位于当前任务内。', 'invalid_input')
        saved = _load_json(plan_path, 'stage plan')
        fingerprint = hashlib.sha256(str(saved.get('prompt') or '').encode('utf-8')).hexdigest()
        if saved.get('stage') != stage or saved.get('promptHash') != fingerprint or ctx.get('stagePromptHash') != fingerprint:
            raise BridgeError('自定义编排的阶段或提示词指纹不一致。', 'conflict')
        renderer = _import_file(Path(__file__).resolve().parent / 'lib/custom_deck.py', '_dsh_custom_deck')
        deckkit = _import_file(_source_path(root, 'src/html/deckkit.py'), '_dsh_custom_deckkit')
        try:
            html = renderer.render(saved.get('plan'), record,
                                   _load_page_texts(job, ctx.get('pagesPath') or ctx.get('pageTextDir')), deckkit)
        except (ValueError, TypeError, KeyError) as exc:
            raise BridgeError(f'自定义演示自检失败：{exc}', 'render_failed') from None
        name = f"{'design' if stage == 'design' else 'PPT'}_{surname}_{year}.html"
        path = artifacts_dir / name
        reused = _persist_immutable([(path, html.encode('utf-8'))])[str(path)]
        produced.append({'stage':stage,'path':str(path),'sha256':_sha256(html.encode('utf-8')),
                         'bytes':len(html.encode('utf-8')),'reused':reused,'selfcheck':'passed',
                         'customPromptHash':fingerprint,'slideCount':len(saved['plan']['slides'])})
    if "design" in stages and not custom_plan_path:
        html_dir = _source_path(root, "src/html")
        if str(html_dir) not in sys.path:
            sys.path.insert(0, str(html_dir))
        design = _import_file(html_dir / "build_design.py", "_dsh_lit_build_design")
        with contextlib.redirect_stdout(logs), contextlib.redirect_stderr(logs):
            deck = design.build_deck(record, topic="", md_path=str(md_path) if md_path.is_file() else None)
            html = deck.render()
            html = _add_markdown_citation_notice(html, record)
            checks = design.selfcheck(html, deck.design_ctx)
        failures = [name for ok, name, _detail in checks if not ok]
        if failures:
            raise BridgeError("实证设计 HTML 自检失败：" + "; ".join(failures[:8]), "render_failed")
        name = f"design_{surname}_{year}.html"
        path = artifacts_dir / name
        reused = _persist_immutable([(path, html.encode("utf-8"))])[str(path)]
        produced.append({"stage": "design", "path": str(path), "sha256": _sha256(html.encode("utf-8")),
                         "bytes": len(html.encode("utf-8")), "reused": reused, "selfcheck": "passed"})
    if "lecture" in stages and not custom_plan_path:
        html_dir = _source_path(root, "src/html")
        if str(html_dir) not in sys.path:
            sys.path.insert(0, str(html_dir))
        lecture = _import_file(html_dir / "build_lecture.py", "_dsh_lit_build_lecture")
        with contextlib.redirect_stdout(logs), contextlib.redirect_stderr(logs):
            deck = lecture.build_deck(record)
            html = deck.render()
            html = _add_markdown_citation_notice(html, record)
            _report, failures = lecture.selfcheck(html, record, deck.audit)
        if failures:
            raise BridgeError("完整介绍 HTML 自检失败：" + "; ".join(failures[:8]), "render_failed")
        name = f"PPT_{surname}_{year}.html"
        path = artifacts_dir / name
        reused = _persist_immutable([(path, html.encode("utf-8"))])[str(path)]
        produced.append({"stage": "lecture", "path": str(path), "sha256": _sha256(html.encode("utf-8")),
                         "bytes": len(html.encode("utf-8")), "reused": reused, "selfcheck": "passed"})
    if "excel" in stages:
        excel_path = _source_path(root, "src/excel/make_excel.py")
        excel = _import_file(excel_path, "_dsh_lit_make_excel")
        validation_problems = excel.validate(record)
        if validation_problems:
            raise BridgeError("Excel 行校验失败：" + "; ".join(validation_problems[:8]), "render_failed")
        rec = copy.deepcopy(record)
        rec["name"] = str(ctx.get("name") or "")
        rec["filled_at"] = str(ctx.get("date") or _dt.date.today().strftime("%Y.%m.%d"))
        deck_links = _excel_deck_links(record, job, render_receipt, pipeline_stages)
        rec["design_deck"] = deck_links["design"]
        rec["lecture_deck"] = deck_links["lecture"]
        path = artifacts_dir / f"{surname}_{year}_文献积累.xlsx"
        semantic_fingerprint = _sha256(_json_bytes({
            "bridge": VERSION, "stage": "excel", "record": rec,
            "sourceSha256": _sha256(excel_path.read_bytes()), "columns": EXPECTED_COLUMNS,
        }))
        prior = (render_receipt.get("rendered") or {}).get("excel")
        if path.is_file():
            actual_hash = _sha256(path.read_bytes())
            if isinstance(prior, dict) and prior.get("path") == str(path) and prior.get("sha256") and prior["sha256"] != actual_hash:
                raise BridgeError("Excel文件hash与render receipt不一致，拒绝复用或覆盖。", "conflict")
            if (isinstance(prior, dict) and prior.get("inputFingerprint")
                    and prior["inputFingerprint"] != semantic_fingerprint):
                raise BridgeError("Excel输入语义指纹与已有产物回执不同，拒绝复用或覆盖。", "conflict")
            if not _excel_semantics_match(path, excel, rec):
                raise BridgeError(f"输出已存在且Excel单元格内容不同，拒绝覆盖：{path}", "conflict")
            produced.append({"stage": "excel", "path": str(path), "sha256": actual_hash,
                             "bytes": path.stat().st_size, "reused": True,
                             "inputFingerprint": semantic_fingerprint,
                             "columns": EXPECTED_COLUMNS, "selfcheck": "passed"})
        else:
            # Persist intent before generating the XLSX. If the process exits
            # after the atomic artifact write but before its final receipt, a
            # retry can verify the workbook cells and safely adopt it.
            _save_render_entry(job, "excel", {
                "path": str(path), "inputFingerprint": semantic_fingerprint,
                "status": "rendering", "selfcheck": "pending",
            })
            wb = excel.build_workbook([rec])
            columns = wb.active.max_column
            if columns != EXPECTED_COLUMNS:
                raise BridgeError("Excel 实际列数不是21，拒绝报告成功。", "render_failed")
            tmp = io.BytesIO()
            wb.save(tmp)
            data = tmp.getvalue()
            try:
                reused = _persist_immutable([(path, data)])[str(path)]
            except BridgeError as exc:
                # A competing render may have atomically created equivalent
                # cells with different ZIP timestamps. Adopt only after full
                # semantic verification; different content remains a conflict.
                if exc.status != "conflict" or not path.is_file() or not _excel_semantics_match(path, excel, rec):
                    raise
                reused = True
            actual_hash = _sha256(path.read_bytes())
            if not _excel_semantics_match(path, excel, rec):
                raise BridgeError("写入后的Excel内容自检失败。", "render_failed")
            produced.append({"stage": "excel", "path": str(path), "sha256": actual_hash,
                             "bytes": path.stat().st_size, "reused": reused,
                             "inputFingerprint": semantic_fingerprint,
                             "columns": columns, "selfcheck": "passed"})
    if produced:
        for artifact in produced:
            entry = {
                "path": artifact["path"], "sha256": artifact["sha256"],
                "selfcheck": artifact.get("selfcheck"), "renderedAt": _dt.datetime.now(_dt.timezone.utc).isoformat(),
                "status": "verified",
            }
            if artifact.get("inputFingerprint"):
                entry["inputFingerprint"] = artifact["inputFingerprint"]
            if artifact.get('customPromptHash'):
                entry['customPromptHash'] = artifact['customPromptHash']
            _save_render_entry(job, artifact["stage"], entry)
    return {"ok": True, "status": "rendered", "selectedStages": stages, "artifacts": produced}


def _cards_library_home(ctx: dict[str, Any]) -> Path:
    """卡片库正本：调用方传入 > KCARD_HOME > 旧默认值（桌面）。"""
    raw = str(ctx.get("libraryHome") or "").strip() or str(os.environ.get("KCARD_HOME") or "").strip()
    if not raw:
        raw = str(Path.home() / "Desktop" / "知识卡片库")
    library_home = Path(raw).expanduser()
    if not library_home.is_absolute():
        raise BridgeError("libraryHome 必须是绝对路径。", "invalid_input")
    return library_home


def _run_make_cards(ctx: dict[str, Any], *, verify: bool) -> dict[str, Any]:
    """跑制卡脚本：默认出《制卡任务书》，verify=True 时做验收。

    脚本不生成卡片正文；正文由执行代理整篇写，这里只负责查重、出任务书和验收。
    """
    record, root = _load_record(ctx)
    job = _job_dir(ctx, root)
    artifacts_dir = job / "artifacts"
    artifacts_dir.mkdir(parents=True, exist_ok=True)
    script = _source_path(root, "src/cards/make_cards.py")
    if not script.is_file():
        raise BridgeError("找不到 src/cards/make_cards.py", "render_failed")
    library_home = _cards_library_home(ctx)
    extract_tmp = artifacts_dir / "_extract_for_cards.json"
    extract_tmp.write_text(json.dumps(record, ensure_ascii=False), encoding="utf-8")
    report = artifacts_dir / ("制卡报告.md" if verify else "制卡任务书.md")
    # 判定书与快照固定放 artifacts：出任务书时写快照，验收时读同一份基线。
    judgement = artifacts_dir / "制卡前判定.json"
    snapshot = artifacts_dir / "制卡候选快照.json"
    env = dict(os.environ)
    env["KCARD_HOME"] = str(library_home)
    env["PYTHONUTF8"] = "1"
    cmd = [sys.executable, "-X", "utf8", str(script),
           "--extract", str(extract_tmp), "--report", str(report),
           "--judgement", str(judgement), "--snapshot", str(snapshot)]
    if verify:
        cmd.append("--verify")
        skips = artifacts_dir / "制卡跳过.json"
        if skips.is_file():
            cmd += ["--skips", str(skips)]
    proc = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", env=env)
    output = ((proc.stdout or "") + (proc.stderr or "")).strip()
    result: dict[str, Any] = {
        "ok": proc.returncode == 0 and report.is_file(),
        "path": str(report) if report.is_file() else None,
        "libraryHome": str(library_home),
        "judgementPath": str(judgement),
        "snapshotPath": str(snapshot),
        "exit": proc.returncode,
        "stdout": output[-4000:],
    }
    if not result["ok"]:
        result["status"] = "invalid"
        # 让代理看到 make_cards 自己的失败摘要，而不是无关的 Excel 校验文案。
        result["message"] = output[-800:].strip() or "制卡脚本失败，无输出"
    else:
        result["status"] = "rendered" if verify else "ready"
    return result


def _action_card_sheet(ctx: dict[str, Any]) -> dict[str, Any]:
    return _run_make_cards(ctx, verify=False)


def _action_cards_verify(ctx: dict[str, Any]) -> dict[str, Any]:
    return _run_make_cards(ctx, verify=True)


def _make_tsv(record: dict[str, Any], ctx: dict[str, Any], job: Path,
              published: dict[str, str] | None = None) -> tuple[list[str], list[str], dict[str, Any]]:
    root = _path(ctx.get("workflowRoot"), "workflowRoot", must_exist=True)
    wechat = _import_file(_source_path(root, "src/wechat/write_tsv.py"), "_dsh_lit_write_tsv")
    rec = copy.deepcopy(record)
    published = published or {}
    design_files = sorted((job / "artifacts").glob("design_*.html")) if (job / "artifacts").is_dir() else []
    lecture_files = sorted((job / "artifacts").glob("PPT_*.html")) if (job / "artifacts").is_dir() else []
    rec["design_url"] = published.get("design") or (design_files[0].name if design_files else "待补")
    rec["ppt_url"] = published.get("lecture") or (lecture_files[0].name if lecture_files else "待补")
    name = str(ctx.get("name") or "")
    today = _dt.date.today()
    date_str = str(ctx.get("date") or f"{today.year}/{today.month}/{today.day}")
    cells, warnings = wechat.build_row(rec, name, date_str, " / ")
    wechat.assert_columns(cells, None)
    registration = {
        "bridge": VERSION,
        "extractPath": str(_path(ctx.get("extractPath"), "extractPath")),
        "paper": (record.get("meta") or {}).get("title") if isinstance(record.get("meta"), dict) else None,
        "author": record.get("authors"), "year": record.get("year"),
        "columns": [x[1] for x in wechat.COLUMN_MAP], "row": cells,
        "columnCount": len(cells), "warnings": warnings,
        "designUrl": rec["design_url"], "lectureUrl": rec["ppt_url"],
    }
    return cells, warnings, registration


def _persist_publish_bundle(job: Path, cells: list[str], registration: dict[str, Any]) -> tuple[Path, Path, str]:
    """Write immutable, content-addressed TSV/registration files.

    Blog publication can add public URLs between a preview and the later sheet
    card. Content-addressed filenames preserve the earlier preview and avoid
    overwriting it when those URLs change the row.
    """
    pub_dir = job / "publish"
    tsv = ("\t".join(cells) + "\n").encode("utf-8")
    digest = _sha256(tsv)
    tsv_path = pub_dir / f"row-{digest[:12]}.tsv"
    registration_path = pub_dir / f"registration-{digest[:12]}.json"
    registration = copy.deepcopy(registration)
    registration["tsvPath"] = str(tsv_path)
    registration["tsvSha256"] = digest
    registration["registrationPath"] = str(registration_path)
    _persist_immutable([(tsv_path, tsv), (registration_path, _json_bytes(registration))])
    return tsv_path, registration_path, digest


def _preview_blog(ctx: dict[str, Any], job: Path, root: Path) -> dict[str, Any]:
    render_receipt, _ = _load_render_receipt(job)
    artifacts_dir = job / "artifacts"
    targets = []
    for stage in ("design", "lecture"):
        path = _verified_rendered_path(render_receipt, artifacts_dir, stage)
        if path is not None:
            targets.append({"stage": stage, "path": str(path), "name": path.name,
                            "sha256": _sha256(path.read_bytes())})
    if not targets:
        return {"ok": False, "status": "needs_setup", "stage": "blog",
                "message": "没有已通过自检且可验证的HTML；请先运行design或lecture卡片。"}

    folder = str(ctx.get("blogFolder") or DEFAULT_BLOG_FOLDER)
    if folder != DEFAULT_BLOG_FOLDER and not ctx.get("allowOtherBlogFolder"):
        return {"ok": False, "status": "needs_setup", "stage": "blog",
                "message": "非 master_fin 分组需显式 allowOtherBlogFolder。"}

    names = [item["name"] for item in targets]
    collisions = sorted({name for name in names if names.count(name) > 1})
    receipt_path = _publish_receipt_path(ctx, job)
    receipt = _receipt_load(receipt_path)
    receipt_conflicts = []
    unresolved = []
    for item in targets:
        key = f"blog:{folder}/{item['name']}"
        prior = receipt["entries"].get(key)
        if prior is not None and not isinstance(prior, dict):
            return {"ok": False, "status": "conflict", "stage": "blog",
                    "message": "博客发布回执条目格式异常；未访问外网。"}
        item["targetFolder"] = folder
        item["targetName"] = item["name"]
        item["alreadyVerified"] = bool(prior and prior.get("status") == "verified" and prior.get("sha256") == item["sha256"])
        if prior and prior.get("sha256") and prior["sha256"] != item["sha256"]:
            receipt_conflicts.append(item["name"])
        elif prior and prior.get("status") in {"prepared", "uploading", "uncertain"}:
            unresolved.append(item["name"])
        if item["alreadyVerified"] and prior.get("url"):
            item["previousUrl"] = prior["url"]
    if collisions or receipt_conflicts:
        return {"ok": False, "status": "conflict", "stage": "blog",
                "message": "本地博客目标名或已有发布回执冲突；未访问外网。",
                "localNameCollisions": collisions, "receiptConflicts": sorted(set(receipt_conflicts))}
    if unresolved:
        return {"ok": False, "status": "uncertain", "stage": "blog",
                "message": "已有未确认的博客发布回执；预览停止以避免重复发布。",
                "unresolvedNames": sorted(set(unresolved))}

    manifest = {
        "bridge": VERSION, "stage": "blog", "targetFolder": folder,
        "files": targets,
        "localNameCollisions": collisions,
        "receiptConflicts": sorted(set(receipt_conflicts)),
        "remoteNamePreflight": "not_performed_no_network",
    }
    digest = _sha256(_json_bytes(manifest))
    preview_path = job / "publish" / f"blog-preview-{digest[:12]}.json"
    manifest["previewPath"] = str(preview_path)
    _persist_immutable([(preview_path, _json_bytes(manifest))])
    return {"ok": True, "status": "preview_ready", "stage": "blog",
            **manifest, "sha256": digest}


def _action_publish_preview(ctx: dict[str, Any]) -> dict[str, Any]:
    record, root = _load_record(ctx)
    job = _job_dir(ctx, root)
    stages = [stage for stage in _stage_list(ctx) if stage in PUBLISH_STAGES]
    if not stages:
        return {"ok": False, "status": "needs_setup", "message": "publish-preview需要显式选择blog或sheet卡。"}
    if len(stages) != 1:
        raise BridgeError("publish-preview每次只接受一个stage：blog或sheet。", "invalid_input")
    validation = validate_record(record, _load_page_texts(job, ctx.get("pagesPath") or ctx.get("pageTextDir")), root)
    if not validation["ok"]:
        return {"ok": False, "status": "invalid", "validation": validation}
    if stages[0] == "blog":
        return _preview_blog(ctx, job, root)

    cells, warnings, registration = _make_tsv(record, ctx, job, _receipt_blog_urls(job, _publish_receipt_path(ctx, job)))
    tsv_path, registration_path, digest = _persist_publish_bundle(job, cells, registration)
    return {"ok": True, "status": "preview_ready", "stage": "sheet", "tsvPath": str(tsv_path),
            "registrationPath": str(registration_path), "columnCount": len(cells),
            "warnings": warnings, "sha256": digest}


def _receipt_load(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {"bridge": VERSION, "entries": {}}
    obj = _load_json(path, "publish receipt")
    if obj.get("bridge") != VERSION or not isinstance(obj.get("entries"), dict):
        raise BridgeError("已有发布回执格式不属于本桥接器，拒绝覆盖。", "conflict")
    return obj


def _publish_receipt_path(ctx: dict[str, Any], job: Path) -> Path:
    candidate = ctx.get("receiptPath")
    if candidate:
        path = _path(candidate, "receiptPath")
        if not _within(path, job):
            raise BridgeError("receiptPath 必须位于 jobDir 内。", "needs_setup")
        return path
    return job / "publish-receipt.json"


def _receipt_blog_urls(job: Path, path: Path | None = None) -> dict[str, str]:
    path = path or job / "publish-receipt.json"
    if not path.is_file():
        return {}
    try:
        receipt = _load_json(path, "publish receipt")
    except BridgeError:
        return {}
    result: dict[str, str] = {}
    for entry in (receipt.get("entries") or {}).values():
        if not isinstance(entry, dict) or entry.get("status") != "verified" or not entry.get("url"):
            continue
        filename = str(entry.get("file") or "")
        if filename.startswith("design_"):
            result["design"] = str(entry["url"])
        elif filename.startswith("PPT_"):
            result["lecture"] = str(entry["url"])
    return result


def _publish_blog(ctx: dict[str, Any], record: dict[str, Any], job: Path,
                  receipt: dict[str, Any], receipt_path: Path) -> dict[str, Any]:
    auth = ctx.get("auth") if isinstance(ctx.get("auth"), dict) else {}
    account, password = auth.get("blogAccount"), auth.get("blogPassword")
    if not (isinstance(account, str) and account and isinstance(password, str) and password):
        return {"ok": False, "status": "needs_auth", "stage": "blog", "message": "缺少博客账号或密码（请通过stdin auth传入）。"}
    root = _path(ctx.get("workflowRoot"), "workflowRoot", must_exist=True)
    upload_mod = _import_file(_source_path(root, "src/blog/upload.py"), "_dsh_lit_blog_upload")
    artifacts_dir = job / "artifacts"
    render_receipt_path = job / "render-receipt.json"
    if not render_receipt_path.is_file():
        return {"ok": False, "status": "needs_setup", "stage": "blog", "message": "当前job没有render receipt；先运行并自检design/lecture卡片。"}
    render_receipt = _load_json(render_receipt_path, "render receipt")
    if render_receipt.get("bridge") != VERSION or not isinstance(render_receipt.get("rendered"), dict):
        return {"ok": False, "status": "conflict", "stage": "blog", "message": "render receipt格式不属于本桥接器或不可用。"}
    targets = []
    for stage in ("design", "lecture"):
        info = (render_receipt.get("rendered") or {}).get(stage)
        if not isinstance(info, dict) or not info.get("path"):
            continue
        path = Path(info["path"]).resolve()
        if not _within(path, artifacts_dir.resolve()) or not path.is_file():
            return {"ok": False, "status": "conflict", "stage": "blog", "message": f"render回执路径无效或文件缺失：{stage}"}
        if _sha256(path.read_bytes()) != info.get("sha256"):
            return {"ok": False, "status": "conflict", "stage": "blog", "message": f"渲染文件hash与回执不一致：{path.name}"}
        targets.append(path)
    if not targets:
        return {"ok": False, "status": "needs_setup", "stage": "blog", "message": "已选择博客发布，但没有对应已渲染HTML；先勾选design/lecture并render。"}
    cfg = _load_json(root / "config.json", "config.json") if (root / "config.json").is_file() else {}
    base = str((cfg.get("blog") or {}).get("baseUrl") or upload_mod.DEFAULT_BASE)
    folder = str(ctx.get("blogFolder") or DEFAULT_BLOG_FOLDER)
    if folder != DEFAULT_BLOG_FOLDER and not ctx.get("allowOtherBlogFolder"):
        return {"ok": False, "status": "needs_setup", "stage": "blog", "message": "非 master_fin 分组需显式 allowOtherBlogFolder。"}
    client = upload_mod.BlogClient(base, account, password)
    try:
        client.login()
        if not client.check_auth():
            return {"ok": False, "status": "needs_auth", "stage": "blog", "message": "博客登录态未确认。"}
        groups = client.get_groups()
        if folder not in groups:
            return {"ok": False, "status": "needs_setup", "stage": "blog", "message": "目标博客分组不存在。"}
        import requests
        results = []
        pages = [p for p in client.list_pages() if p.get("group") == folder]
        for path in targets:
            local = path.read_bytes()
            digest = _sha256(local)
            name = path.name
            key = f"blog:{folder}/{name}"
            prior = receipt["entries"].get(key) if isinstance(receipt.get("entries"), dict) else None
            remote = next((p for p in pages if p.get("filename") == name), None)
            if remote:
                url = client.view_url(folder, name)
                try:
                    resp = requests.get(url, timeout=30, allow_redirects=True)
                except requests.RequestException:
                    return {"ok": False, "status": "uncertain", "stage": "blog", "message": f"无法回读已有文件：{name}"}
                remote_hash = _sha256(resp.content) if resp.status_code == 200 else None
                if remote_hash == digest:
                    receipt["entries"][key] = {"status": "verified", "file": name, "group": folder,
                                                "sha256": digest, "url": url,
                                                "verifiedAt": _dt.datetime.now(_dt.timezone.utc).isoformat(),
                                                "recoveredByReadback": True}
                    _atomic_json(receipt_path, receipt, owned_marker=VERSION)
                    results.append({"stage": "blog", "file": name, "status": "verified", "url": url, "recovered": True})
                    continue
                return {"ok": False, "status": "conflict", "stage": "blog", "message": f"远端已有同名文件但内容hash不同：{name}；不覆盖、不重发。"}
            if prior and prior.get("status") in {"prepared", "uploading", "uncertain"}:
                return {"ok": False, "status": "uncertain", "stage": "blog", "message": f"上次发布结果未确认且当前列表未显示目标：{name}；为避免重复上传已停止。"}
            # Durable intent receipt before the irreversible upload.
            receipt["entries"][key] = {"status": "prepared", "file": name, "group": folder,
                                        "sha256": digest, "preparedAt": _dt.datetime.now(_dt.timezone.utc).isoformat()}
            _atomic_json(receipt_path, receipt, owned_marker=VERSION)
            try:
                client.upload(path, name, folder)
                receipt["entries"][key]["status"] = "uploading"
                receipt["entries"][key]["uploadedAt"] = _dt.datetime.now(_dt.timezone.utc).isoformat()
                _atomic_json(receipt_path, receipt, owned_marker=VERSION)
                after_pages = [p for p in client.list_pages() if p.get("group") == folder and p.get("filename") == name]
                if len(after_pages) != 1:
                    receipt["entries"][key]["status"] = "uncertain"
                    _atomic_json(receipt_path, receipt, owned_marker=VERSION)
                    return {"ok": False, "status": "uncertain", "stage": "blog", "message": f"上传后未能唯一回查：{name}。"}
                url = client.view_url(folder, name)
                resp = requests.get(url, timeout=30, allow_redirects=True)
                remote_hash = _sha256(resp.content) if resp.status_code == 200 else None
                if remote_hash != digest:
                    receipt["entries"][key]["status"] = "uncertain"
                    receipt["entries"][key]["url"] = url
                    _atomic_json(receipt_path, receipt, owned_marker=VERSION)
                    return {"ok": False, "status": "uncertain", "stage": "blog", "message": f"上传已回查但公开内容hash未匹配：{name}。"}
                receipt["entries"][key].update({"status": "verified", "url": url,
                                               "verifiedAt": _dt.datetime.now(_dt.timezone.utc).isoformat()})
                _atomic_json(receipt_path, receipt, owned_marker=VERSION)
                results.append({"stage": "blog", "file": name, "status": "verified", "url": url, "recovered": False})
                pages.extend(after_pages)
            except Exception:
                receipt["entries"][key]["status"] = "uncertain"
                _atomic_json(receipt_path, receipt, owned_marker=VERSION)
                return {"ok": False, "status": "uncertain", "stage": "blog", "message": f"博客发布异常，回执已保留且不会盲目重发：{name}。"}
        return {"ok": True, "status": "published", "stage": "blog", "files": results}
    except Exception as exc:
        return {"ok": False, "status": "error", "stage": "blog", "message": f"博客预检失败：{type(exc).__name__}。"}


def _sheet_tool_args(tool: dict[str, Any], values: dict[str, Any]) -> dict[str, Any] | None:
    schema = tool.get("inputSchema") or {}
    props = schema.get("properties") or {}
    required = schema.get("required") or []
    aliases = {
        "keyword": ("keyword", "query", "search_key", "searchKey", "text"),
        "file_id": ("file_id", "fileId", "node_id", "nodeId"),
        "sheet_name": ("sheet_name", "sheetName", "sheet"),
        "sheet_id": ("sheet_id", "sheetId"),
        "values": ("values", "data", "rows"),
        "range": ("range", "sheet_range", "sheetRange", "a1_range", "a1Range"),
        "start_row": ("start_row", "startRow"),
        "end_row": ("end_row", "endRow"),
        "start_col": ("start_col", "startCol"),
        "end_col": ("end_col", "endCol"),
        "return_csv": ("return_csv", "returnCsv"),
        "include_formula": ("include_formula", "includeFormula"),
    }
    result = {}
    for semantic, names in aliases.items():
        if semantic not in values or values[semantic] is None:
            continue
        target = next((name for name in names if name in props), None)
        if target:
            result[target] = values[semantic]
    if any(k not in result for k in required):
        return None
    return result


def _sheet_write_strategy(tool: dict[str, Any]) -> str | None:
    """Fail closed unless the MCP description guarantees a server-side append.

    A range update can overwrite a row created by another editor after our
    read. Version/CAS is intentionally unsupported until the concrete MCP
    schema and read response expose a matching atomic precondition.
    """
    schema = tool.get("inputSchema") or {}
    required = set(schema.get("required") or [])
    range_names = {"range", "sheet_range", "sheetRange", "a1_range", "a1Range"}
    if required.intersection(range_names):
        return None
    description = str(tool.get("description") or "").casefold()
    english_append = re.search(
        r"\bappend(?:s|ed|ing)?\s+(?:(?:a|the|new)\s+)?(?:row|rows|data|values)\b"
        r".{0,100}\b(?:to|at)\s+(?:the\s+)?(?:end|bottom)\b",
        description,
    )
    chinese_append = re.search(r"(?:追加|写入).{0,16}(?:到|至)(?:表格|工作表)?(?:末尾|最后)", description)
    chinese_append_reverse = re.search(r"(?:表格|工作表)?(?:末尾|最后).{0,12}追加", description)
    if english_append or chinese_append or chinese_append_reverse:
        return "append"
    return None


def _mcp_json(result: Any) -> Any:
    if isinstance(result, dict) and result.get("isError"):
        raise BridgeError("腾讯文档 MCP 工具返回错误。", "needs_setup")
    if isinstance(result, dict):
        for item in result.get("content") or []:
            if isinstance(item, dict) and item.get("type") == "text":
                try:
                    return json.loads(item.get("text") or "")
                except json.JSONDecodeError:
                    continue
        if isinstance(result.get("structuredContent"), (dict, list)):
            return result["structuredContent"]
    return result


def _find_bound_document_id(value: Any, document_url: str) -> str | None:
    """Only bind IDs on a result object carrying the configured document URL."""
    def identity(raw: Any) -> tuple[str, str] | None:
        if not isinstance(raw, str):
            return None
        try:
            parsed = urlsplit(raw)
            if parsed.scheme != "https" or parsed.hostname != "docs.qq.com" or parsed.port not in (None, 443):
                return None
            return (parsed.hostname, parsed.path.rstrip("/"))
        except ValueError:
            return None

    target = identity(document_url)
    if not target or not target[1].startswith("/sheet/"):
        return None
    matched: set[str] = set()

    def visit(item: Any) -> None:
        if isinstance(item, dict):
            urls = (item.get(k) for k in ("url", "file_url", "fileUrl", "doc_url", "docUrl", "web_url", "webUrl", "link"))
            if any(identity(url) == target for url in urls):
                matched.update(str(item[k]) for k in ("node_id", "file_id", "fileId", "nodeId") if item.get(k))
            for child in item.values():
                visit(child)
        elif isinstance(item, list):
            for child in item:
                visit(child)

    visit(value)
    return next(iter(matched)) if len(matched) == 1 else None


def _extract_rows(payload: Any) -> list[list[Any]] | None:
    if isinstance(payload, list) and all(isinstance(r, (list, tuple)) for r in payload):
        return [list(r) for r in payload]
    if isinstance(payload, dict):
        for key in ("values", "data", "rows", "cells"):
            child = payload.get(key)
            if isinstance(child, list):
                if child and isinstance(child[0], dict) and key == "data":
                    continue
                if all(isinstance(r, (list, tuple)) for r in child):
                    return [list(r) for r in child]
        for child in payload.values():
            found = _extract_rows(child)
            if found is not None:
                return found
    return None


def _clean_cell(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _normalise_rows(rows: list[list[Any]], width: int = EXPECTED_COLUMNS) -> list[list[str]]:
    if any(len(r) > width for r in rows):
        raise BridgeError("腾讯文档返回超过21列的数据，拒绝写入。", "needs_setup")
    out = [[_clean_cell(v) for v in r] + [""] * (width - len(r)) for r in rows]
    while out and not any(out[-1]):
        out.pop()
    return out


_SHEET_READ_CHUNK_ROWS = 40


@contextlib.contextmanager
def _modern_sheet_lock(job: Path, file_id: str, sheet_id: str):
    """Serialize this plugin's writers for one document/tab across processes.

    The one-byte OS lock is released automatically when a process exits. The
    lock file is intentionally retained; stale locks are never guessed away.
    """
    key = _sha256((file_id + "\0" + sheet_id).encode("utf-8"))
    lock_dir = job.parent / ".sheet-locks"
    lock_dir.mkdir(parents=True, exist_ok=True)
    lock_path = lock_dir / (key + ".lock")
    handle = lock_path.open("a+b")
    try:
        handle.seek(0, os.SEEK_END)
        if handle.tell() == 0:
            handle.write(b"\0")
            handle.flush()
        handle.seek(0)
        if os.name == "nt":
            import msvcrt
            msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
        else:
            import fcntl
            fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
    except OSError as exc:
        handle.close()
        raise BridgeError("另一进程正在登记此工作表；本次未写入。", "busy") from exc
    try:
        yield key
    finally:
        try:
            handle.seek(0)
            if os.name == "nt":
                import msvcrt
                msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                import fcntl
                fcntl.flock(handle.fileno(), fcntl.LOCK_UN)
        finally:
            handle.close()


def _modern_tool_call(client: Any, tool: dict[str, Any], args: dict[str, Any]) -> Any:
    with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
        payload = _mcp_json(client.call_tool(tool["name"], args))
    if isinstance(payload, dict) and payload.get("error") not in (None, ""):
        raise BridgeError("腾讯表格读取工具返回错误。", "needs_setup")
    return payload


def _modern_sheet_dimensions(client: Any, info_tool: dict[str, Any], file_id: str,
                             sheet_id: str, sheet_name: str) -> tuple[int, int]:
    args = _sheet_tool_args(info_tool, {"file_id": file_id})
    if args is None:
        raise BridgeError("新版工作表信息参数与 schema 不匹配。", "needs_setup")
    payload = _modern_tool_call(client, info_tool, args)
    if not isinstance(payload, dict) or not isinstance(payload.get("sheets"), list):
        raise BridgeError("新版接口未返回可核对的工作表清单。", "needs_setup")
    matches = [s for s in payload["sheets"] if isinstance(s, dict)
               and s.get("sheet_id") == sheet_id and s.get("sheet_name") == sheet_name]
    if len(matches) != 1:
        raise BridgeError("读取期间工作表身份发生变化或不再唯一。", "conflict")
    sheet = matches[0]
    row_count = sheet.get("row_count")
    column_count = sheet.get("column_count", sheet.get("col_count"))
    if (not isinstance(row_count, int) or isinstance(row_count, bool) or row_count < 1
            or not isinstance(column_count, int) or isinstance(column_count, bool) or column_count < EXPECTED_COLUMNS):
        raise BridgeError("新版接口未提供完整、有效的工作表行列容量。", "needs_setup")
    return row_count, column_count


def _modern_cell_value(cell: dict[str, Any]) -> dict[str, Any]:
    kind = str(cell.get("value_type") or "").upper()
    if kind not in {"STRING", "NUMBER", "BOOL", "FORMULA"}:
        raise BridgeError("工作表回读包含未知单元格类型；未写入。", "needs_setup")
    value: dict[str, Any] = {"value_type": kind}
    formula = cell.get("formula")
    if isinstance(formula, str) and formula:
        value["formula"] = formula
    if kind == "STRING":
        raw = cell.get("string_value")
        if not isinstance(raw, str):
            raise BridgeError("工作表字符串单元格缺少 string_value。", "needs_setup")
        value["string_value"] = raw
    elif kind == "NUMBER":
        raw = cell.get("number_value")
        if not isinstance(raw, (int, float)) or isinstance(raw, bool):
            raise BridgeError("工作表数值单元格缺少 number_value。", "needs_setup")
        value["number_value"] = raw
    elif kind == "BOOL":
        raw = cell.get("bool_value")
        if not isinstance(raw, bool):
            raise BridgeError("工作表布尔单元格缺少 bool_value。", "needs_setup")
        value["bool_value"] = raw
    else:
        raw_formula = cell.get("formula")
        if not isinstance(raw_formula, str) or not raw_formula:
            raise BridgeError("工作表公式单元格缺少 formula。", "needs_setup")
        value["formula"] = raw_formula
        # Formula reads include their calculated typed value as well. Retain
        # whichever result fields the service returned so later diffs catch
        # recalculation, while ignoring the schema's unused default fields.
        for key, expected_type in (("string_value", str), ("number_value", (int, float)),
                                   ("bool_value", bool)):
            if key not in cell:
                continue
            raw = cell[key]
            if expected_type == (int, float):
                valid = isinstance(raw, (int, float)) and not isinstance(raw, bool)
            else:
                valid = isinstance(raw, expected_type)
            if valid:
                value[key] = raw
    return value


def _modern_cell_nonempty(value: dict[str, Any]) -> bool:
    if value.get("formula"):
        return True
    if value["value_type"] == "STRING":
        return value.get("string_value", "") != ""
    return True


def _modern_read_range(client: Any, cell_tool: dict[str, Any], file_id: str,
                       sheet_id: str, start_row: int, end_row: int,
                       start_col: int, end_col: int) -> dict[tuple[int, int], dict[str, Any]]:
    args = _sheet_tool_args(cell_tool, {
        "file_id": file_id, "sheet_id": sheet_id,
        "start_row": start_row, "end_row": end_row,
        "start_col": start_col, "end_col": end_col,
        "return_csv": False, "include_formula": True,
    })
    requested = {"file_id", "sheet_id", "start_row", "end_row", "start_col", "end_col",
                 "return_csv", "include_formula"}
    if args is None or not requested.issubset(args):
        raise BridgeError("新版单元格读取 schema 无法表达指定的完整区域。", "needs_setup")
    payload = _modern_tool_call(client, cell_tool, args)
    if not isinstance(payload, dict) or not isinstance(payload.get("cells"), list):
        raise BridgeError("新版单元格读取没有返回可验证的 cells 数组。", "needs_setup")
    result: dict[tuple[int, int], dict[str, Any]] = {}
    for raw in payload["cells"]:
        if not isinstance(raw, dict):
            raise BridgeError("新版单元格读取包含无效元素。", "needs_setup")
        row, col = raw.get("row"), raw.get("col")
        if (not isinstance(row, int) or isinstance(row, bool) or not isinstance(col, int)
                or isinstance(col, bool) or not (start_row <= row <= end_row)
                or not (start_col <= col <= end_col)):
            raise BridgeError("新版单元格读取返回了请求范围之外的坐标。", "needs_setup")
        coord = (row, col)
        if coord in result:
            raise BridgeError("新版单元格读取重复返回同一坐标。", "needs_setup")
        value = _modern_cell_value(raw)
        # The service returns sparse cells. An empty string is semantically
        # blank; formula cells remain occupied even when they calculate to "".
        if _modern_cell_nonempty(value):
            result[coord] = value
    return result


def _modern_read_snapshot(client: Any, info_tool: dict[str, Any], cell_tool: dict[str, Any],
                          file_id: str, sheet_id: str, sheet_name: str) -> dict[str, Any]:
    before_dims = _modern_sheet_dimensions(client, info_tool, file_id, sheet_id, sheet_name)
    row_count, column_count = before_dims
    cells: dict[tuple[int, int], dict[str, Any]] = {}
    for start_row in range(0, row_count, _SHEET_READ_CHUNK_ROWS):
        end_row = min(start_row + _SHEET_READ_CHUNK_ROWS, row_count) - 1
        block = _modern_read_range(client, cell_tool, file_id, sheet_id,
                                   start_row, end_row, 0, column_count - 1)
        overlap = set(cells).intersection(block)
        if overlap:
            raise BridgeError("分块读取返回了重叠坐标。", "needs_setup")
        cells.update(block)
    after_dims = _modern_sheet_dimensions(client, info_tool, file_id, sheet_id, sheet_name)
    if after_dims != before_dims:
        raise BridgeError("全表读取期间工作表容量发生变化；未写入。", "conflict")
    return {"rowCount": row_count, "columnCount": column_count, "cells": cells}


def _modern_snapshot_json(snapshot: dict[str, Any]) -> dict[str, Any]:
    cells = [{"row": row, "col": col, **value}
             for (row, col), value in sorted(snapshot["cells"].items())]
    return {"rowCount": snapshot["rowCount"], "columnCount": snapshot["columnCount"], "cells": cells}


def _modern_snapshot_from_json(value: Any) -> dict[str, Any] | None:
    if not isinstance(value, dict) or not isinstance(value.get("cells"), list):
        return None
    row_count, column_count = value.get("rowCount"), value.get("columnCount")
    if not isinstance(row_count, int) or not isinstance(column_count, int):
        return None
    cells: dict[tuple[int, int], dict[str, Any]] = {}
    try:
        for raw in value["cells"]:
            row, col = raw.get("row"), raw.get("col")
            if not isinstance(row, int) or not isinstance(col, int):
                return None
            coord = (row, col)
            if coord in cells:
                return None
            cells[coord] = _modern_cell_value(raw)
    except (AttributeError, BridgeError):
        return None
    return {"rowCount": row_count, "columnCount": column_count, "cells": cells}


def _modern_snapshots_equal(left: dict[str, Any], right: dict[str, Any]) -> bool:
    return (left["rowCount"] == right["rowCount"]
            and left["columnCount"] == right["columnCount"]
            and left["cells"] == right["cells"])


def _modern_row_is_empty(snapshot: dict[str, Any], row: int, end_col: int) -> bool:
    return not any(_modern_cell_nonempty(value)
                   for (cell_row, col), value in snapshot["cells"].items()
                   if cell_row == row and 0 <= col <= end_col)


def _modern_row_matches(snapshot: dict[str, Any], row: int, cells: list[str]) -> bool:
    for col, expected in enumerate(cells):
        actual = snapshot["cells"].get((row, col))
        if expected == "":
            if actual is not None and _modern_cell_nonempty(actual):
                return False
        elif (actual is None or actual.get("value_type") != "STRING"
              or actual.get("formula") or actual.get("string_value") != expected):
            return False
    return True


def _modern_write_result_matches(snapshot: dict[str, Any], before: dict[str, Any],
                                 row: int, cells: list[str]) -> bool:
    expected = dict(before["cells"])
    for col, value in enumerate(cells):
        if value == "":
            expected.pop((row, col), None)
        else:
            expected[(row, col)] = {"value_type": "STRING", "string_value": value}
    return (snapshot["rowCount"] == before["rowCount"]
            and snapshot["columnCount"] == before["columnCount"]
            and snapshot["cells"] == expected)


def _publish_modern_sheet(ctx: dict[str, Any], cells: list[str], job: Path,
                          receipt: dict[str, Any], receipt_path: Path, client: Any,
                          info_tool: dict[str, Any], cell_tool: dict[str, Any],
                          write_tool: dict[str, Any], file_id: str, sheet_id: str,
                          sheet_name: str, doc_url: str) -> dict[str, Any]:
    if len(cells) != EXPECTED_COLUMNS or any(not isinstance(c, str) or "\n" in c or "\t" in c or "\r" in c for c in cells):
        return {"ok": False, "status": "invalid", "stage": "sheet", "message": "TSV 行不是严格的21个单行单元格。"}
    if not write_tool:
        return {"ok": False, "status": "needs_setup", "stage": "sheet",
                "reasonCode": "modern_range_writer_unavailable",
                "message": "新版表格读取可用，但缺少 set_range_value；未写入。"}
    write_args_base = _sheet_tool_args(write_tool, {"file_id": file_id, "sheet_id": sheet_id,
                                                    "values": []})
    if (write_args_base is None or "values" not in write_args_base
            or not (set(write_args_base) >= {"file_id", "sheet_id", "values"})):
        return {"ok": False, "status": "needs_setup", "stage": "sheet",
                "reasonCode": "modern_range_writer_schema_mismatch",
                "message": "set_range_value schema 未能精确绑定文档、工作表和单元格值；未写入。"}

    expected_row_hash = _sha256(_json_bytes(cells))
    row_values = [{"row": None, "col": col, "value_type": "STRING", "string_value": value}
                  for col, value in enumerate(cells)]
    entries = receipt.setdefault("entries", {})
    identity_hash = _sha256((file_id + "\0" + sheet_id).encode("utf-8"))
    try:
        with _modern_sheet_lock(job, file_id, sheet_id):
            before = _modern_read_snapshot(client, info_tool, cell_tool, file_id, sheet_id, sheet_name)
            row_count = before["rowCount"]
            # Existing intents are always reconciled by reads first. A write
            # that might already have been dispatched is never sent again.
            matching = [(key, entry) for key, entry in entries.items()
                        if isinstance(entry, dict) and entry.get("sheetIdentitySha256") == identity_hash]
            for key, old in matching:
                if old.get("status") == "aborted_before_write":
                    continue
                if old.get("rowSha256") != expected_row_hash:
                    return {"ok": False, "status": "conflict", "stage": "sheet",
                            "message": "此 job 对同一工作表已有不同内容的登记回执；拒绝追加第二行。"}
                old_target = old.get("targetRow")
                if not isinstance(old_target, int) or old_target < 1:
                    return {"ok": False, "status": "uncertain", "stage": "sheet",
                            "message": "已有工作表登记回执缺少有效目标行；不会再次写入。"}
                target0 = old_target - 1
                if old.get("status") == "deduplicated":
                    if _modern_row_matches(before, target0, cells):
                        return {"ok": True, "status": "published", "stage": "sheet",
                                "recovered": True, "deduplicated": True,
                                "targetRow": old_target, "rowSha256": expected_row_hash}
                    return {"ok": False, "status": "conflict", "stage": "sheet",
                            "message": "已登记的重复行回读不一致；拒绝再次写入。"}
                before_digest = old.get("beforeSnapshotSha256")
                if not isinstance(before_digest, str) or not re.fullmatch(r"[0-9a-f]{64}", before_digest):
                    return {"ok": False, "status": "uncertain", "stage": "sheet",
                            "message": "已有写入回执缺少有效的基线 hash；不会再次写入。"}
                audit_path = job / ".sheet-append-audit" / (
                    f"{identity_hash[:16]}-{old_target}-{expected_row_hash[:16]}-{before_digest[:16]}.json")
                try:
                    audit = _load_json(audit_path, "sheet append audit")
                    audit_snapshot = audit.get("beforeSnapshot")
                    if _sha256(_json_bytes(audit_snapshot)) != before_digest:
                        audit_snapshot = None
                    baseline = _modern_snapshot_from_json(audit_snapshot)
                except (BridgeError, OSError):
                    baseline = None
                if baseline is None:
                    return {"ok": False, "status": "uncertain", "stage": "sheet",
                            "message": "已有写入意图缺少可验证的原始快照；不会再次写入。"}
                if _modern_write_result_matches(before, baseline, target0, cells):
                    recovered_status = "deduplicated" if old.get("status") == "prepared" else "verified"
                    old.update({"status": recovered_status,
                                "afterSnapshotSha256": _sha256(_json_bytes(_modern_snapshot_json(before))),
                                "verifiedAt": old.get("verifiedAt") or _dt.datetime.now(_dt.timezone.utc).isoformat()})
                    try:
                        _atomic_json(receipt_path, receipt, owned_marker=VERSION)
                    except Exception:
                        return {"ok": False, "status": "uncertain", "stage": "sheet",
                                "message": "表格已回读到预期行，但回执保存失败；下次只会读取核对。"}
                    return {"ok": True, "status": "published", "stage": "sheet", "recovered": True,
                            "deduplicated": recovered_status == "deduplicated",
                            "targetRow": old_target, "rowSha256": expected_row_hash}
                if old.get("status") == "prepared":
                    # This status is persisted before the separate inflight
                    # marker, so no write request has been dispatched yet.
                    old["status"] = "aborted_before_write"
                    old["abortedAt"] = _dt.datetime.now(_dt.timezone.utc).isoformat()
                    _atomic_json(receipt_path, receipt, owned_marker=VERSION)
                    continue
                if old.get("status") in {"inflight", "uncertain", "writing"}:
                    if _modern_snapshots_equal(before, baseline):
                        return {"ok": False, "status": "uncertain", "stage": "sheet",
                                "targetRow": old_target,
                                "message": "已有登记意图尚未出现在表中；为避免重复写入，本次只读核对并停止。"}
                    return {"ok": False, "status": "conflict", "stage": "sheet",
                            "targetRow": old_target,
                            "message": "已有登记意图与当前表格不一致；不覆盖、不重发。"}
                return {"ok": False, "status": "conflict", "stage": "sheet",
                        "message": "已有工作表登记回执状态无法安全恢复；拒绝再次写入。"}

            # Older receipts lack a sheet identity hash. Block unresolved
            # attempts for the same intended row rather than risk a duplicate.
            for key, old in entries.items():
                if (isinstance(old, dict) and key.startswith("sheet:")
                        and old.get("rowSha256") == expected_row_hash
                        and old.get("status") in {"prepared", "writing", "inflight", "uncertain"}):
                    return {"ok": False, "status": "uncertain", "stage": "sheet",
                            "message": "旧版回执记录了尚未核清的同一行；本次只读后停止，不会重发。"}
                if (isinstance(old, dict) and key.startswith("sheet:")
                        and "sheetIdentitySha256" not in old and old.get("status") == "verified"):
                    return {"ok": False, "status": "conflict", "stage": "sheet",
                            "message": "旧版回执已有已验证的工作表行；请使用新 job 登记不同论文。"}

            # Deduplicate exact 21-cell strings before choosing a new row.
            duplicates = [row for row in range(row_count)
                          if _modern_row_matches(before, row, cells)]
            if duplicates:
                target_row = duplicates[0] + 1
                key = f"sheet:{identity_hash}:{target_row}"
                entries[key] = {"status": "deduplicated", "targetRow": target_row,
                                "sheetIdentitySha256": identity_hash,
                                "rowSha256": expected_row_hash,
                                "verifiedAt": _dt.datetime.now(_dt.timezone.utc).isoformat()}
                _atomic_json(receipt_path, receipt, owned_marker=VERSION)
                return {"ok": True, "status": "published", "stage": "sheet", "recovered": True,
                        "deduplicated": True, "targetRow": target_row, "rowSha256": expected_row_hash}

            # SAFETY.md reserves row 1 for the teacher's header. An empty or
            # malformed sheet must never cause the first data row to land there.
            if _modern_row_is_empty(before, 0, EXPECTED_COLUMNS - 1):
                return {"ok": False, "status": "needs_setup", "stage": "sheet",
                        "reasonCode": "sheet_header_missing",
                        "message": "第1行 A:U 未检测到表头；为保护老师表头，未写入。"}

            last_data_row0 = max((row for (row, _col), value in before["cells"].items()
                                  if _modern_cell_nonempty(value)), default=-1)
            target0 = last_data_row0 + 1
            target_row = target0 + 1
            if target0 >= row_count:
                return {"ok": False, "status": "needs_setup", "stage": "sheet",
                        "reasonCode": "sheet_capacity_reached",
                        "message": "表格容量内没有末行之后的空行；未写入。"}
            if not _modern_row_is_empty(before, target0, EXPECTED_COLUMNS - 1):
                return {"ok": False, "status": "conflict", "stage": "sheet", "targetRow": target_row,
                        "message": "目标行 A:U 含有内容，已停止写入。"}
            # Since the last row was computed over every column, the target
            # must also be empty beyond U before any A:U write.
            if any(row == target0 and _modern_cell_nonempty(value)
                   for (row, _col), value in before["cells"].items()):
                return {"ok": False, "status": "conflict", "stage": "sheet", "targetRow": target_row,
                        "message": "目标行存在 A:U 之外的内容，已停止写入。"}

            audit_dir = job / ".sheet-append-audit"
            before_json = _modern_snapshot_json(before)
            before_digest = _sha256(_json_bytes(before_json))
            audit_path = audit_dir / (
                f"{identity_hash[:16]}-{target_row}-{expected_row_hash[:16]}-{before_digest[:16]}.json")
            audit_value = {"bridge": VERSION, "fileId": file_id, "sheetId": sheet_id,
                           "sheetName": sheet_name, "targetRow": target_row,
                           "rowSha256": expected_row_hash, "beforeSnapshot": before_json,
                           "beforeSnapshotSha256": before_digest}
            _persist_immutable([(audit_path, _json_bytes(audit_value))])
            receipt_key = f"sheet:{identity_hash}:{target_row}:{before_digest}"
            intent = {"status": "prepared", "targetRow": target_row,
                      "sheetIdentitySha256": identity_hash, "rowSha256": expected_row_hash,
                      "beforeSnapshotSha256": before_digest,
                      "auditFile": audit_path.name,
                      "preparedAt": _dt.datetime.now(_dt.timezone.utc).isoformat()}
            entries[receipt_key] = intent
            _atomic_json(receipt_path, receipt, owned_marker=VERSION)

            # Full-grid re-read catches edits after the baseline was captured.
            precommit = _modern_read_snapshot(client, info_tool, cell_tool, file_id, sheet_id, sheet_name)
            if not _modern_snapshots_equal(before, precommit):
                return {"ok": False, "status": "conflict", "stage": "sheet", "targetRow": target_row,
                        "message": "全表在预检后发生变化；保留 prepared 意图，本次未写入。"}
            if not _modern_row_is_empty(precommit, target0, EXPECTED_COLUMNS - 1):
                return {"ok": False, "status": "conflict", "stage": "sheet", "targetRow": target_row,
                        "message": "提交前目标行 A:U 已出现内容；本次未写入。"}

            # Narrow last-moment target check, directly before the sole write.
            target_cells = _modern_read_range(client, cell_tool, file_id, sheet_id,
                                              target0, target0, 0, EXPECTED_COLUMNS - 1)
            if target_cells:
                return {"ok": False, "status": "conflict", "stage": "sheet", "targetRow": target_row,
                        "message": "最终检查发现目标行 A:U 已非空；本次未写入。"}
            intent["status"] = "inflight"
            intent["writeStartedAt"] = _dt.datetime.now(_dt.timezone.utc).isoformat()
            _atomic_json(receipt_path, receipt, owned_marker=VERSION)

            for item in row_values:
                item["row"] = target0
            write_args = dict(write_args_base)
            write_args["values"] = row_values
            write_error = False
            try:
                _modern_tool_call(client, write_tool, write_args)
            except Exception:
                # The server may have committed even when the client saw an
                # error. Always reconcile by reading; never repeat this call.
                write_error = True
            try:
                after = _modern_read_snapshot(client, info_tool, cell_tool, file_id, sheet_id, sheet_name)
            except Exception:
                intent["status"] = "uncertain"
                intent["writeResultUnknown"] = True
                _atomic_json(receipt_path, receipt, owned_marker=VERSION)
                return {"ok": False, "status": "uncertain", "stage": "sheet", "targetRow": target_row,
                        "message": "写入已尝试但写后回读失败；保留意图，之后只读核对，不会重发。"}
            if _modern_write_result_matches(after, before, target0, cells):
                intent.update({"status": "verified", "afterSnapshotSha256": _sha256(_json_bytes(_modern_snapshot_json(after))),
                               "verifiedAt": _dt.datetime.now(_dt.timezone.utc).isoformat(),
                               "writeAcknowledged": not write_error})
                try:
                    _atomic_json(receipt_path, receipt, owned_marker=VERSION)
                except Exception:
                    return {"ok": False, "status": "uncertain", "stage": "sheet", "targetRow": target_row,
                            "message": "表格内容已核对正确，但回执保存失败；下次只会读取核对。"}
                return {"ok": True, "status": "published", "stage": "sheet", "recovered": False,
                        "targetRow": target_row, "rowSha256": expected_row_hash,
                        "writeAcknowledged": not write_error,
                        "atomicAppend": False}
            if _modern_snapshots_equal(after, before):
                intent["status"] = "uncertain"
                intent["afterSnapshotSha256"] = _sha256(_json_bytes(_modern_snapshot_json(after)))
                _atomic_json(receipt_path, receipt, owned_marker=VERSION)
                return {"ok": False, "status": "uncertain", "stage": "sheet", "targetRow": target_row,
                        "message": "写入尝试后表格仍与原快照相同；不会重发，后续只读核对。"}
            intent["status"] = "uncertain"
            intent["afterSnapshotSha256"] = _sha256(_json_bytes(_modern_snapshot_json(after)))
            _atomic_json(receipt_path, receipt, owned_marker=VERSION)
            return {"ok": False, "status": "conflict", "stage": "sheet", "targetRow": target_row,
                    "message": "写后完整回读与预期不一致；未修复、不覆盖，也不会重发。"}
    except BridgeError as exc:
        return {"ok": False, "status": exc.status, "stage": "sheet", "message": str(exc)}
    except Exception as exc:
        return {"ok": False, "status": "needs_setup", "stage": "sheet",
                "message": f"新版表格安全预检失败：{type(exc).__name__}；未确认写入。"}


def _publish_sheet(ctx: dict[str, Any], cells: list[str], job: Path,
                   receipt: dict[str, Any], receipt_path: Path) -> dict[str, Any]:
    auth = ctx.get("auth") if isinstance(ctx.get("auth"), dict) else {}
    token = auth.get("tencentToken")
    if not isinstance(token, str) or not token.strip():
        return {"ok": False, "status": "needs_auth", "stage": "sheet", "message": "缺少腾讯文档 MCP token（请通过stdin auth传入）。"}
    root = _path(ctx.get("workflowRoot"), "workflowRoot", must_exist=True)
    log = _import_file(_source_path(root, "src/wechat/log_sheet.py"), "_dsh_lit_log_sheet")
    endpoint = str(ctx.get("mcpEndpoint") or log.MCP_ENDPOINT)
    doc_title = str(ctx.get("docTitle") or log.DOC_TITLE)
    doc_url = str(ctx.get("docUrl") or getattr(log, "DOC_URL", "https://docs.qq.com/sheet/DTHRhWHFwaGJPTWtL"))
    sheet_name = str(ctx.get("sheetName") or "2026年9月")
    transport = _import_file(Path(__file__).resolve().parent / 'lib/tencent_mcp.py', '_dsh_lit_tencent_mcp')
    client = transport.McpClient(endpoint, token)
    # Official Skill short-code tokens may use the raw Authorization value.
    # Keep older OAuth Bearer tokens compatible; never modify the original workflow client.
    if auth.get("tencentAuthScheme") == "raw":
        original_headers = client._headers
        def raw_headers():
            headers = original_headers()
            headers["Authorization"] = token
            return headers
        client._headers = raw_headers
    try:
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            client.initialize()
            tools = client.list_tools()
        by_name = {t.get("name"): t for t in tools if isinstance(t, dict)}
        search, read, write = (by_name.get("manage.search_file") or by_name.get("search_space_file"), by_name.get("get_content"),
                               by_name.get("batch_update_sheet_range"))
        modern = by_name.get('sheet.get_sheet_info') and by_name.get('sheet.get_cell_data')
        if not search or (not modern and not (read and write)):
            return {"ok": False, "status": "needs_setup", "stage": "sheet", "message": "MCP 缺少可识别的搜索及表格读取工具，未写入。"}
        search_args = _sheet_tool_args(search, {"keyword": doc_title})
        if search_args is None:
            return {"ok": False, "status": "needs_setup", "stage": "sheet", "message": "目标表搜索参数与已知 schema 不匹配。"}
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            found = _mcp_json(client.call_tool(search['name'], search_args))
        file_id = str(ctx.get("nodeId") or _find_bound_document_id(found, doc_url) or "")
        if not file_id:
            return {"ok": False, "status": "needs_setup", "stage": "sheet", "message": "搜索结果未唯一绑定已配置的腾讯文档链接；请配置明确的 nodeId，未读取或写入其他文档。"}
        if modern:
            info_tool = by_name['sheet.get_sheet_info']
            info_args = _sheet_tool_args(info_tool, {'file_id':file_id})
            if info_args is None:
                raise BridgeError('新版工作表信息参数未能完整匹配，未写入。', 'needs_setup')
            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                info = _mcp_json(client.call_tool(info_tool['name'], info_args))
            if not isinstance(info, dict) or info.get('error') or not isinstance(info.get('sheets'), list):
                raise BridgeError('新版接口未返回可核对的工作表清单，未写入。', 'needs_setup')
            requested_id = ctx.get('sheetId')
            matched = [s for s in info['sheets'] if isinstance(s, dict) and s.get('sheet_name') == sheet_name
                       and (not requested_id or s.get('sheet_id') == requested_id)]
            if len(matched) != 1 or not matched[0].get('sheet_id'):
                raise BridgeError('未唯一匹配指定工作表名称与ID，未写入其他工作表。', 'needs_setup')
            selected = matched[0]
            result = _publish_modern_sheet(ctx, cells, job, receipt, receipt_path, client,
                                           info_tool, by_name['sheet.get_cell_data'],
                                           by_name.get('sheet.set_range_value'), file_id,
                                           str(selected['sheet_id']), sheet_name, doc_url)
            result.setdefault('target', {'documentURL':doc_url,'fileId':file_id,
                                         'sheetId':selected['sheet_id'],'sheetName':sheet_name})
            result.setdefault('capabilities', {'modernDiscovery':True,'readCells':True,
                                               'guardedRangeWrite':bool(by_name.get('sheet.set_range_value')),
                                               'atomicAppend':False})
            return result
        read_args = _sheet_tool_args(read, {"file_id": file_id, "sheet_name": sheet_name,
                                            "sheet_id": ctx.get("sheetId")})
        if read_args is None:
            return {"ok": False, "status": "needs_setup", "stage": "sheet", "message": "get_content 必填参数超出已知映射，需人工适配。"}
        target_keys = {"sheet_name", "sheetName", "sheet", "sheet_id", "sheetId"}
        file_keys = {"file_id", "fileId", "node_id", "nodeId"}
        if not target_keys.intersection(read_args) or not file_keys.intersection(read_args):
            return {"ok": False, "status": "needs_setup", "stage": "sheet", "message": "get_content 未绑定明确文档与工作表，未写入。"}
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            before_payload = _mcp_json(client.call_tool("get_content", read_args))
        before_raw = _extract_rows(before_payload)
        if before_raw is None:
            return {"ok": False, "status": "needs_setup", "stage": "sheet", "message": "get_content 未返回可验证的完整二维表格；未写入。"}
        before = _normalise_rows(before_raw)
        if len(cells) != EXPECTED_COLUMNS or any("\n" in c or "\t" in c or "\r" in c for c in cells):
            return {"ok": False, "status": "invalid", "stage": "sheet", "message": "TSV 行不是严格的21个单行单元格。"}
        expected_row_hash = _sha256(_json_bytes([_clean_cell(x) for x in cells]))
        # On retry, inspect old receipt rows before deriving a new append position.
        # This prevents an already-committed row from being appended a second time.
        for old_key, old in receipt.get("entries", {}).items():
            if not old_key.startswith("sheet:") or not isinstance(old, dict):
                continue
            old_target = old.get("targetRow")
            if not isinstance(old_target, int) or old_target < 1:
                continue
            current_row = before[old_target - 1] if old_target - 1 < len(before) else None
            prefix = before[:old_target - 1]
            if (current_row and _sha256(_json_bytes(current_row)) == old.get("rowSha256")
                    and _sha256(_json_bytes(prefix)) == old.get("prefixSha256")):
                if old.get("rowSha256") == expected_row_hash:
                    old["status"] = "verified"
                    old["afterRowsSha256"] = _sha256(_json_bytes(before))
                    old["verifiedAt"] = old.get("verifiedAt") or _dt.datetime.now(_dt.timezone.utc).isoformat()
                    _atomic_json(receipt_path, receipt, owned_marker=VERSION)
                    return {"ok": True, "status": "published", "stage": "sheet", "recovered": True,
                            "targetRow": old_target, "rowSha256": old.get("rowSha256")}
                if old.get("status") == "verified":
                    return {"ok": False, "status": "conflict", "stage": "sheet", "message": "此job已有已验证的不同TSV行；拒绝重复登记。"}
            if old.get("status") in {"prepared", "writing", "uncertain"}:
                return {"ok": False, "status": "uncertain", "stage": "sheet", "message": "已有未完成腾讯表回执；已读取但不会自动重发，请人工核对。"}
            if old.get("status") == "verified":
                return {"ok": False, "status": "conflict", "stage": "sheet", "message": "已验证的历史行与当前回读不一致；拒绝再次写入。"}
        last_row = max((i + 1 for i, row in enumerate(before) if any(row)), default=0)
        target_row = last_row + 1
        if target_row - 1 < len(before) and any(before[target_row - 1]):
            return {"ok": False, "status": "conflict", "stage": "sheet", "message": "目标行 A:U 非空，已停止。"}
        before_hash = _sha256(_json_bytes(before))
        receipt_key = "sheet:" + str(target_row)
        write_strategy = _sheet_write_strategy(write)
        if write_strategy != "append":
            return {"ok": False, "status": "needs_setup", "stage": "sheet",
                    "message": "写工具未明确保证服务端末尾追加；范围更新在共享表并发下可能覆盖他人内容，未写入。"}
        write_args = _sheet_tool_args(write, {"file_id": file_id, "sheet_name": sheet_name,
                                              "sheet_id": ctx.get("sheetId"),
                                              "values": [cells]})
        if write_args is None:
            return {"ok": False, "status": "needs_setup", "stage": "sheet", "message": "写入工具必填参数超出已知映射，未写入。"}
        if not target_keys.intersection(write_args) or not file_keys.intersection(write_args):
            return {"ok": False, "status": "needs_setup", "stage": "sheet", "message": "写入工具未绑定明确文档与工作表，未写入。"}
        # Persist intent before the external side effect, without storing existing sheet content or credentials.
        receipt["entries"][receipt_key] = {
            "status": "prepared", "targetRow": target_row, "beforeRowsSha256": before_hash,
            "prefixSha256": _sha256(_json_bytes(before[:target_row - 1])),
            "rowSha256": expected_row_hash, "preparedAt": _dt.datetime.now(_dt.timezone.utc).isoformat(),
        }
        _atomic_json(receipt_path, receipt, owned_marker=VERSION)
        try:
            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                client.call_tool("batch_update_sheet_range", write_args)
                after_payload = _mcp_json(client.call_tool("get_content", read_args))
            after_raw = _extract_rows(after_payload)
            if after_raw is None:
                receipt["entries"][receipt_key]["status"] = "uncertain"
                _atomic_json(receipt_path, receipt, owned_marker=VERSION)
                return {"ok": False, "status": "uncertain", "stage": "sheet", "message": "写后回读失败；保留回执，不会盲目重发。"}
            after = _normalise_rows(after_raw)
            expected = list(before)
            while len(expected) < target_row - 1:
                expected.append([""] * EXPECTED_COLUMNS)
            if len(expected) == target_row - 1:
                expected.append([_clean_cell(x) for x in cells])
            else:
                expected[target_row - 1] = [_clean_cell(x) for x in cells]
            if after != expected:
                receipt["entries"][receipt_key]["status"] = "uncertain"
                receipt["entries"][receipt_key]["afterRowsSha256"] = _sha256(_json_bytes(after))
                _atomic_json(receipt_path, receipt, owned_marker=VERSION)
                return {"ok": False, "status": "conflict", "stage": "sheet", "message": "写后全表与预期逐格不一致；不尝试修复或重写。"}
            receipt["entries"][receipt_key].update({"status": "verified", "afterRowsSha256": _sha256(_json_bytes(after)),
                                                     "verifiedAt": _dt.datetime.now(_dt.timezone.utc).isoformat()})
            _atomic_json(receipt_path, receipt, owned_marker=VERSION)
            return {"ok": True, "status": "published", "stage": "sheet", "recovered": False,
                    "targetRow": target_row, "rowSha256": expected_row_hash}
        except Exception:
            receipt["entries"][receipt_key]["status"] = "uncertain"
            _atomic_json(receipt_path, receipt, owned_marker=VERSION)
            return {"ok": False, "status": "uncertain", "stage": "sheet", "message": "腾讯表写入/回读异常；保留回执，不会盲目重发。"}
    except BridgeError as exc:
        return {"ok": False, "status": exc.status, "stage": "sheet", "message": str(exc)}
    except Exception as exc:
        return {"ok": False, "status": "needs_setup", "stage": "sheet", "message": f"腾讯文档预检失败：{type(exc).__name__}。"}


def _action_publish(ctx: dict[str, Any]) -> dict[str, Any]:
    record, root = _load_record(ctx)
    job = _job_dir(ctx, root)
    stages = [s for s in _stage_list(ctx) if s in PUBLISH_STAGES]
    if not stages:
        return {"ok": False, "status": "needs_setup", "message": "publish 需要显式选择单个 stage：blog 或 sheet。", "stages": []}
    if len(stages) != 1:
        raise BridgeError("publish 每次只接受一个 stage：blog 或 sheet。", "invalid_input")
    validation = validate_record(record, _load_page_texts(job, ctx.get("pagesPath") or ctx.get("pageTextDir")), root)
    if not validation["ok"]:
        return {"ok": False, "status": "invalid", "validation": validation}
    pub_dir = job / "publish"
    receipt_path = _publish_receipt_path(ctx, job)
    receipt = _receipt_load(receipt_path)
    stage_results: list[dict[str, Any]] = []
    published_urls: dict[str, str] = _receipt_blog_urls(job, receipt_path)
    if "blog" in stages:
        result = _publish_blog(ctx, record, job, receipt, receipt_path)
        stage_results.append(result)
        if result.get("ok"):
            for item in result.get("files", []):
                if item.get("file", "").startswith("design_"):
                    published_urls["design"] = item.get("url", "")
                elif item.get("file", "").startswith("PPT_"):
                    published_urls["lecture"] = item.get("url", "")
    if "sheet" in stages:
        published_urls.update(_receipt_blog_urls(job, receipt_path))
        cells, warnings, registration = _make_tsv(record, ctx, job, published_urls)
        tsv_path, registration_path, row_hash = _persist_publish_bundle(job, cells, registration)
        sheet_result = _publish_sheet(ctx, cells, job, receipt, receipt_path)
        sheet_result.update({'tsvPath':str(tsv_path),'registrationPath':str(registration_path),
                             'rowSha256':sheet_result.get('rowSha256') or row_hash,'columnCount':EXPECTED_COLUMNS})
        stage_results.append(sheet_result)
    states = [r.get("status") for r in stage_results]
    ok = all(r.get("ok") for r in stage_results)
    status = "completed" if ok else ("partial" if any(r.get("ok") for r in stage_results) else (states[0] if len(set(states)) == 1 else "blocked"))
    message = next((str(r.get("message")) for r in stage_results if not r.get("ok") and r.get("message")), None)
    return {"ok": ok, "status": status, "selectedStages": stages, "message": message,
            **{key:stage_results[0][key] for key in ('tsvPath','registrationPath') if key in stage_results[0]},
            "stages": stage_results, "receiptPath": str(receipt_path) if receipt_path.exists() else None}


def dispatch(action: str, ctx: dict[str, Any]) -> dict[str, Any]:
    if not isinstance(ctx, dict):
        raise BridgeError("stdin 必须是 JSON 对象。", "invalid_input")
    actions = {"probe": _action_probe, "prepare": _action_prepare, "validate": _action_validate,
               "render": _render, "card-sheet": _action_card_sheet, "cards-verify": _action_cards_verify,
               "publish-preview": _action_publish_preview, "publish": _action_publish}
    if action not in actions:
        raise BridgeError("action 仅支持 probe / prepare / validate / render / card-sheet / cards-verify / publish-preview / publish。", "invalid_input")
    return actions[action](ctx)


def _redact_result(value: Any, ctx: dict[str, Any]) -> Any:
    secrets: list[str] = []
    auth = ctx.get("auth") if isinstance(ctx.get("auth"), dict) else {}
    for key, val in auth.items():
        if str(key).casefold() in _SECRET_KEYS and isinstance(val, str) and len(val) >= 4:
            secrets.append(val)
    def clean(x: Any) -> Any:
        if isinstance(x, dict):
            return {k: clean(v) for k, v in x.items() if str(k).casefold() not in _SECRET_KEYS}
        if isinstance(x, list):
            return [clean(v) for v in x]
        if isinstance(x, str):
            for secret in secrets:
                x = x.replace(secret, "[REDACTED]")
            return x
        return x
    return clean(value)


def main(argv: list[str] | None = None) -> int:
    args = list(sys.argv[1:] if argv is None else argv)
    action = args[0] if len(args) == 1 else ""
    try:
        raw = sys.stdin.read()
        ctx = json.loads(raw) if raw.strip() else {}
        result = dispatch(action, ctx)
        result = _redact_result(result, ctx if isinstance(ctx, dict) else {})
        status = result.get("status", "ok")
        result.setdefault("ok", status in {"ready", "valid", "rendered", "preview_ready", "published", "skipped"})
        sys.stdout.write(json.dumps(result, ensure_ascii=False, separators=(",", ":")) + "\n")
        return 0 if result.get("ok") else 1
    except BridgeError as exc:
        result = {"ok": False, "status": exc.status, "error": str(exc)}
        sys.stdout.write(json.dumps(result, ensure_ascii=False, separators=(",", ":")) + "\n")
        return 1
    except Exception as exc:  # do not print tracebacks; exception text can include external payloads
        result = {"ok": False, "status": "error", "error": f"处理失败：{type(exc).__name__}。"}
        sys.stdout.write(json.dumps(result, ensure_ascii=False, separators=(",", ":")) + "\n")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
